#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""「分组管理（增删改查）+ 导入面板分组列」E2E 探针。

★ 本探针跑在**本机模式**：main() 把注入的 window.WB_CFG 正则为空 → WBCloud/WBGrp.off() 为 true，
  所有写入退回 localStorage，wb_c100_series / wb_c100_groups 两张表一条都不动。
  （分组注册表也走同一套降级：WBGrp.off() 时 GRP_STATE='off'，只写 wb_c100_groups_v1 缓存。）

覆盖：
  0. 页面形态：本机模式、GRP 起点为空、管理面板里能看到「分组管理」入口
  1. 导入面板「分组」列：表头、每行一个下拉、默认值全是「我的商品」、默认组排在选项第一位
  2. 分组管理面板：打开（走真实按钮）、「我的商品」固定行不可改名、列结构
  3. 增：新建空分组「钼系」「贵金属」→ 出现在列表里、商品 0、显示「暂无商品」
  4. 空分组已能被导入面板选到（下拉里出现）
  5. 逐行设分组导入（钼系 / 贵金属 / 我的商品）→ 三条商品的 tag 与预期一致
  6. 查：分组管理里的商品数 / 数据点 / 区间随导入更新
  7. 改：把「钼系」改名「钼系金属」→ 注册表与组内商品一起改名
  8. 合并：改名撞上已存在的「有色金属」→ 两组并一组，旧名消失
  9. 删：删除「有色金属」→ 组内商品退回「我的商品」且数据不丢
 10. 套用到所有行：默认分组一次盖到全部行 + 自由文本新组名自动建档
 11. 备份 .json 含 groups；localStorage 缓存与内存 GRP 一致
 12. 全程零 JS 错误

跑法：
  /Users/luoshui/.workbuddy/binaries/python/envs/default/bin/python probe_grp_e2e.py
"""
import functools
import http.server
import io
import json
import os
import re
import shutil
import socketserver
import subprocess
import sys
import threading
import time

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
SRC_PAGE = os.path.join(ROOT, 'deliver', '大宗商品日线对比图-2005-2026.html')
SITE = os.path.join(HERE, '_grp_site')            # 独立临时目录：绝不碰 deploy_site/
PAGE = 'commodity-daily.html'
PROBE_PAGE = '_e2e_grp.html'
XLSX_NAME = '_grp_src.xlsx'
XLSX = os.path.join(SITE, XLSX_NAME)
PORT = 8807
EDGE = '/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge'
RESULT = []


def make_xlsx():
    """一列日期 + 3 个数值列，够跑逐行分组与批量套用两组用例"""
    import datetime
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = '日线'
    ws.append(['日期', '黄金 · COMEX', '白银 · COMEX', '钼铁 · FeMo60'])
    ws.append(['单位', '美元/盎司', '美元/盎司', '元/吨'])
    d = datetime.date(2024, 1, 1)
    i = 0
    while i < 60:
        if d.weekday() < 5:
            ws.append([d, round(2000 + i * 3.3, 2), round(24 + i * 0.11, 3), round(138000 + i * 220.0, 1)])
        d += datetime.timedelta(days=1)
        i += 1
    for r in range(3, ws.max_row + 1):
        ws.cell(row=r, column=1).number_format = 'yyyy-mm-dd'
    wb.save(XLSX)
    return os.path.getsize(XLSX), ws.max_row


DBG = r'''
<pre id="dbg" style="display:none"></pre>
<script>
(function(){
  var errs = [], out = [];
  window.addEventListener('error', function(e){ errs.push('ERR: ' + e.message); });
  window.addEventListener('unhandledrejection', function(e){
    errs.push('REJ: ' + ((e.reason && e.reason.message) || e.reason));
  });
  window.confirm = function(){ return true; };
  window.prompt = function(){ return ''; };

  function log(k, v){ out.push(k + ' = ' + v); }
  function sec(t){ out.push(''); out.push('=== ' + t + ' ==='); }
  function id(i){ return document.getElementById(i); }
  function txt(el){ return el ? (el.textContent || '').replace(/\s+/g, ' ').trim() : '(null)'; }
  function click(el){ if(!el){ out.push('✗ 找不到元素'); return false; } el.click(); return true; }
  function clickId(i){ return click(id(i)); }
  function sleep(ms){ return new Promise(function(r){ setTimeout(r, ms); }); }
  function waitFor(fn, ms, tag){
    return new Promise(function(res, rej){
      var t0 = Date.now();
      (function poll(){
        var v; try{ v = fn(); }catch(e){ return rej(new Error(tag + ' 抛错: ' + e.message)); }
        if(v) return res(v);
        if(Date.now() - t0 > (ms || 3000)) return rej(new Error(tag + ' 超时(' + (ms || 3000) + 'ms)'));
        requestAnimationFrame(poll);
      })();
    });
  }
  function setVal(el, v){
    if(!el) return false;
    el.value = v;
    el.dispatchEvent(new Event('input',  { bubbles: true }));
    el.dispatchEvent(new Event('change', { bubbles: true }));
    return true;
  }
  function qa(sel){ return Array.prototype.slice.call(document.querySelectorAll(sel)); }

  /* ---- 分组相关的读取助手 ---- */
  function tagCounts(){
    var m = {};
    CUSTOM.filter(function(s){ return s.src !== 'overlay'; }).forEach(function(s){
      var t = tagOf(s); m[t] = (m[t] || 0) + 1;
    });
    return Object.keys(m).sort().map(function(k){ return k + ' ' + m[k]; }).join(' · ');
  }
  function grpNames(){ return GRP.map(function(g){ return g.n; }).join(' / ') || '(空)'; }
  function grpRowOf(name){
    var rows = qa('#dlg table.ftb tr');
    for(var i = 0; i < rows.length; i++){
      var inp = rows[i].querySelector('input[data-gn]');
      if(inp && inp.dataset.gn === name) return rows[i];
    }
    return null;
  }
  function panelRow(name){
    /* 「我的商品」那行没有 input，用文字匹配兜底 */
    var rows = qa('#dlg table.ftb tr');
    for(var i = 0; i < rows.length; i++){
      var inp = rows[i].querySelector('input[data-gn]');
      var key = inp ? inp.dataset.gn : txt(rows[i].querySelector('td'));
      if(key === name || key.indexOf(name + ' 默认分组') === 0) return rows[i];
    }
    return null;
  }
  function panelLine(name){
    var r = panelRow(name);
    if(!r) return '(找不到该分组行)';
    return Array.prototype.map.call(r.querySelectorAll('td'), function(td){
      var i = td.querySelector('input[data-gn]');
      return (i ? i.value : txt(td)).replace(/\s+/g, ' ').trim();
    }).join(' | ');
  }
  function panelNames(){
    var got = [];
    qa('#dlg table.ftb tr').forEach(function(r){
      var i = r.querySelector('input[data-gn]');
      if(i){ got.push(i.dataset.gn); return; }
      var bb = r.querySelector('td > b');           // 「我的商品」那行的名字是 <b>
      if(bb){ got.push(txt(bb)); return; }
      if(r.querySelector('button[data-gdel]')) got.push(txt(r.querySelector('td')));
    });
    return got.join(' / ') || '(无数据行)';
  }
  function openGrp(){
    clickId('btn-my');
    if(!id('mg-grp')) return false;
    clickId('mg-grp');
    return !!id('gm-add');
  }
  function newGrp(name){
    if(!id('gm-new')) return false;
    setVal(id('gm-new'), name);
    clickId('gm-add');
    return true;
  }
  function renameGrp(oldName, newName){
    var r = grpRowOf(oldName);
    if(!r) return false;
    setVal(r.querySelector('input[data-gn]'), newName);
    var b = r.querySelector('button[data-gsave]');
    if(!b || b.disabled) return false;
    b.click();
    return true;
  }
  function delGrp(name){
    var r = grpRowOf(name);
    if(!r) return false;
    var b = r.querySelector('button[data-gdel]');
    if(!b) return false;
    b.click();
    return true;
  }
  async function openUpAndParse(){
    clickId('btn-up');
    await waitFor(function(){ return !!id('fpick'); }, 5000, '打开上传面板');
    var ab = await (await fetch('@XLSX@')).arrayBuffer();
    readUpFile(new File([ab], '分组测试源.xlsx'));
    await waitFor(function(){ return UP.list && UP.list.length >= 3; }, 20000, '解析 xlsx');
  }
  function rowTagSelect(colIdx){
    return document.querySelector('#up-list select[data-tg="' + colIdx + '"]');
  }
  function setRowTag(colIdx, tag){
    var s = rowTagSelect(colIdx);
    if(!s) return false;
    var has = Array.prototype.some.call(s.options, function(o){ return o.value === tag; });
    if(!has) return false;                       // 选项必须由页面自己给出，探针不硬塞
    s.value = tag;
    s.dispatchEvent(new Event('change', { bubbles: true }));
    return true;
  }

  (async function(){
    try{
      /* ============ 0. 页面形态 ============ */
      sec('0. 页面形态（本机模式）');
      log('window.WB_CFG（应为 null）', JSON.stringify(window.WB_CFG || null));
      log('WBCloud.off()', WBCloud.off());
      log('WBGrp.off()', WBGrp.off());
      log('云表名', GROUP_TABLE);
      log('GRP 起点条数', GRP.length);
      log('GRP_STATE 起点', GRP_STATE);
      log('localStorage 键', GRP_KEY);
      log('HAS_BUILTIN', HAS_BUILTIN);

      /* ============ 1. 导入面板的「分组」列 ============ */
      sec('1. 导入面板：分组列（默认「我的商品」）');
      await openUpAndParse();
      log('预览表列数', qa('#up-list tr').length ? qa('#up-list tr')[0].querySelectorAll('th').length : 0);
      log('表头', Array.prototype.map.call(qa('#up-list tr')[0].querySelectorAll('th'),
            function(t){ return txt(t); }).join(' | '));
      log('行数 / 分组下拉数', UP.list.length + ' / ' + qa('#up-list select[data-tg]').length);
      log('每行下拉默认值不同者（应 0）',
          qa('#up-list select[data-tg]').filter(function(s){ return s.value !== '我的商品'; }).length);
      var s0 = qa('#up-list select[data-tg]')[0];
      log('下拉第一个选项（应是默认组）', s0 ? s0.options[0].value : '(无)');
      log('下拉选项', s0 ? Array.prototype.map.call(s0.options, function(o){ return o.value; }).join(' / ') : '(无)');
      log('「默认分组」输入框在', !!id('up-tag'));
      log('「套用到所有行」按钮在', !!id('up-tagall'));
      closeDlg();

      /* ============ 2. 分组管理面板（入口 + 结构） ============ */
      sec('2. 分组管理：从管理面板进入');
      log('打开分组管理', openGrp());
      log('面板标题', txt(document.querySelector('#dlg h3')));
      log('表头', Array.prototype.map.call(qa('#dlg table.ftb tr')[0].querySelectorAll('th'),
            function(t){ return txt(t); }).join(' | '));
      log('列出的分组', panelNames());
      log('「我的商品」行原文', panelLine('我的商品'));
      var dflt = panelRow('我的商品');
      log('默认行没有改名输入框（应 true）', !(dflt && dflt.querySelector('input[data-gn]')));
      log('默认行没有删除按钮（应 true）', !(dflt && dflt.querySelector('button[data-gdel]')));
      log('新建输入框 maxlength', id('gm-new') ? id('gm-new').getAttribute('maxlength') : '(无)');

      /* ============ 3. 增：新建两个空分组 ============ */
      sec('3. 增：新建空分组');
      newGrp('钼系');
      log('新建后 GRP', grpNames());
      log('列表里的分组', panelNames());
      log('「钼系」行原文（商品数应为 0）', panelLine('钼系'));
      newGrp('贵金属');
      log('再建后 GRP', grpNames());
      log('列表里的分组', panelNames());
      log('空分组提示是否出现', txt(panelRow('钼系')).indexOf('暂无商品') >= 0);
      /* 重名 / 空名保护 */
      var n0 = GRP.length;
      newGrp('钼系');
      log('重复新建同名不新增（应 true）', GRP.length === n0);
      newGrp('我的商品');
      log('拿默认组名新建不新增（应 true）', GRP.length === n0);
      newGrp('');
      log('空名字不新增（应 true）', GRP.length === n0);

      /* ============ 4. 空分组能被导入面板选到 ============ */
      sec('4. 空分组出现在导入面板下拉里');
      closeDlg();
      await openUpAndParse();
      var opts = qa('#up-list select[data-tg]')[0];
      var vals = opts ? Array.prototype.map.call(opts.options, function(o){ return o.value; }) : [];
      log('下拉选项', vals.join(' / '));
      log('含「钼系」（应 true）', vals.indexOf('钼系') >= 0);
      log('含「贵金属」（应 true）', vals.indexOf('贵金属') >= 0);
      log('默认组仍在第一位（应 true）', vals[0] === '我的商品');

      /* ============ 5. 逐行设分组后导入 ============ */
      sec('5. 逐行设分组 → 导入');
      var cs = UP.list.map(function(it){ return it.c; });
      log('第 1 行设为「钼系」', setRowTag(cs[0], '钼系'));
      log('第 2 行设为「贵金属」', setRowTag(cs[1], '贵金属'));
      log('第 3 行保持默认', rowTagSelect(cs[2]) ? rowTagSelect(cs[2]).value : '(无)');
      log('已选统计', txt(id('up-selinfo')));
      log('导入前提示', txt(id('up-stat')).slice(0, 120));
      var n0 = CUSTOM.length;
      clickId('up-ok');
      await waitFor(function(){ return CUSTOM.length > n0; }, 6000, '导入落地');
      log('新增条数', CUSTOM.length - n0);
      log('CUSTOM 分组分布', tagCounts());
      log('GRP（应仍为 贵金属 / 钼系）', grpNames());

      /* ============ 6. 查：统计随导入更新 ============ */
      sec('6. 查：分组统计');
      log('打开分组管理', openGrp());
      log('列出的分组', panelNames());
      log('「我的商品」行', panelLine('我的商品'));
      log('「钼系」行', panelLine('钼系'));
      log('「贵金属」行', panelLine('贵金属'));
      log('统计卡文字', txt(document.querySelector('#dlg h4 .n')));
      var g1 = groupStats();
      log('groupStats 顺序', g1.map(function(g){ return g.name + '(' + g.n + ')'; }).join(' / '));
      log('默认组排在最后（应 true）', g1[g1.length - 1].name === '我的商品');
      var mg = g1.filter(function(g){ return g.name === '钼系'; })[0];
      log('钼系 点数 > 0（应 true）', !!mg && mg.pts > 0);
      log('钼系 区间', mg ? (dayToStr(mg.t0) + ' → ' + dayToStr(mg.t1)) : '(无)');

      /* ============ 7. 改：重命名 ============ */
      sec('7. 改：重命名分组（商品跟着搬家）');
      log('改名「钼系」→「钼系金属」', renameGrp('钼系', '钼系金属'));
      log('GRP', grpNames());
      log('列表里的分组', panelNames());
      log('CUSTOM 分组分布', tagCounts());
      log('「钼系金属」商品数（应 1）', (groupStats().filter(function(g){ return g.name === '钼系金属'; })[0] || {n:-1}).n);
      log('旧名「钼系」已不在 GRP（应 true）', GRP.every(function(g){ return g.n !== '钼系'; }));

      /* ============ 8. 合并：改名撞上已有分组 ============ */
      sec('8. 合并：撞上已有分组');
      newGrp('有色金属');
      log('先建一个空的「有色金属」', GRP.map(function(g){ return g.n; }).join(' / '));
      log('改名「钼系金属」→「有色金属」（＝合并）', renameGrp('钼系金属', '有色金属'));
      log('GRP（应只剩 有色金属 / 贵金属）', grpNames());
      log('列表里的分组', panelNames());
      log('旧名「钼系金属」已消失（应 true）', GRP.every(function(g){ return g.n !== '钼系金属'; }));
      log('「有色金属」商品数（应 1）',
          (groupStats().filter(function(g){ return g.name === '有色金属'; })[0] || {n:-1}).n);

      /* ============ 9. 删：分组删除，商品退回默认组 ============ */
      sec('9. 删：删除分组（商品退回「我的商品」）');
      var before = CUSTOM.length;
      log('删除「有色金属」', delGrp('有色金属'));
      log('GRP', grpNames());
      log('列表里的分组', panelNames());
      log('商品总数不变（数据没丢）', CUSTOM.length + ' (删前 ' + before + ')');
      log('CUSTOM 分组分布（应 我的商品 2 · 贵金属 1）', tagCounts());

      /* ============ 10. 套用到所有行 + 自由文本新组名自动建档 ============ */
      sec('10. 套用到所有行 / 新组名自动建档');
      closeDlg();
      await openUpAndParse();
      setVal(id('up-tag'), '钼系');
      clickId('up-tagall');
      var vals2 = qa('#up-list select[data-tg]').map(function(s){ return s.value; });
      log('套用后各行分组', vals2.join(' / '));
      log('全部套用成功（应 true）', vals2.length > 0 && vals2.every(function(v){ return v === '钼系'; }));
      log('新增了下拉选项「钼系」（应 true）',
          Array.prototype.some.call(qa('#up-list select[data-tg]')[0].options, function(o){ return o.value === '钼系'; }));
      var n1 = CUSTOM.length;
      clickId('up-ok');
      await waitFor(function(){ return CUSTOM.length > n1; }, 6000, '批量导入落地');
      log('新增条数', CUSTOM.length - n1);
      log('CUSTOM 分组分布（应 我的商品 2 · 贵金属 1 · 钼系 3）', tagCounts());
      log('自由文本新组名已登记进 GRP（应 true）', !!grpFind('钼系'));
      log('GRP', grpNames());

      /* ============ 11. 备份 .json 含 groups + 本机缓存一致 ============ */
      sec('11. 备份 / 缓存');
      clickId('btn-bak');
      await waitFor(function(){ return !!id('bk-out'); }, 4000, '打开备份面板');
      var dls = [], real = window.download;
      window.download = function(name, data, mime){ dls.push({ name: name, data: data }); };
      clickId('bk-out');
      await sleep(400);
      window.download = real;
      log('拦到的下载数（应 1）', dls.length);
      log('文件名', dls[0] ? dls[0].name : '(无)');
      var bk = null;
      try{ bk = JSON.parse(dls[0].data); }catch(e){}
      log('JSON 可解析', !!bk);
      log('备份版本号', bk ? bk.v : '(无)');
      log('备份 groups 条数', bk && bk.groups ? bk.groups.length : '(无)');
      log('备份 groups 名字', bk && bk.groups ? bk.groups.map(function(g){ return g.n; }).join(' / ') : '(无)');
      log('备份 series 条数', bk && bk.series ? bk.series.length : '(无)');
      log('备份面板说明提到分组', txt(document.querySelector('#dlg .sub2')).indexOf('分组') >= 0);
      log('备份面板标题行', txt(document.querySelector('#dlg h4')));
      closeDlg();
      var cached = [];
      try{ cached = JSON.parse(localStorage.getItem('wb_c100_groups_v1') || '[]'); }catch(e){}
      log('本机缓存 groups', cached.map(function(g){ return g.n; }).join(' / '));
      log('缓存与内存一致（应 true）',
          cached.length === GRP.length && cached.every(function(g, i){ return g.n === GRP[i].n; }));
      log('缓存字段齐全', cached.length ? Object.keys(cached[0]).sort().join(',') : '(空)');
    }catch(e){
      out.push('✗ 探针中断：' + (e && e.message || e));
    }
    out.push('');
    out.push('JS错误: ' + (errs.length ? errs.join(' | ') : '无'));
    await fetch('/__res__', { method: 'POST', body: out.join('\n') + '\n' });
  })();
})();
</script>
'''


class Quiet(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def do_POST(self):
        n = int(self.headers.get('Content-Length') or 0)
        body = self.rfile.read(n).decode('utf-8', 'ignore')
        RESULT.append(body)
        self.send_response(200)
        self.send_header('Content-Length', '2')
        self.end_headers()
        self.wfile.write(b'ok')


def serve():
    h = functools.partial(Quiet, directory=SITE)
    socketserver.TCPServer.allow_reuse_address = True
    httpd = socketserver.TCPServer(('127.0.0.1', PORT), h)
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    return httpd


def run_probe(url, timeout=240):
    del RESULT[:]
    cmd = [EDGE, '--headless=new', '--no-sandbox', '--disable-gpu', '--disable-dev-shm-usage',
           '--window-size=1512,1600', url]
    p = subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    t0 = time.time()
    while time.time() - t0 < timeout:
        if RESULT:
            break
        time.sleep(0.25)
    time.sleep(0.4)
    try:
        p.terminate(); p.wait(timeout=8)
    except Exception:
        try: p.kill()
        except Exception: pass
    return RESULT[0] if RESULT else '（%.0f 秒内没收到探针回传）' % timeout


def main():
    shutil.rmtree(SITE, ignore_errors=True)
    os.makedirs(SITE, exist_ok=True)
    size, rows = make_xlsx()
    print('=' * 78)
    print('「分组管理（增删改查）+ 导入面板分组列」E2E')
    print('=' * 78)
    print('源 xlsx：%s（%d 字节 · %d 行 · 3 个数值列）' % (XLSX_NAME, size, rows))

    src = io.open(SRC_PAGE, encoding='utf-8').read()
    # ★ 本机模式：掏空云端配置 —— 保证探针一张生产表都不写
    stripped, cnt = re.subn(r'window\.WB_CFG\s*=\s*\{[^}]*\};', 'window.WB_CFG = null;', src)
    assert cnt == 1, '未能掏空 WB_CFG（匹配 %d 次）' % cnt
    i = src.find('<body')
    j = src.find('>', i)
    pre = ('<script>try{localStorage.removeItem("wb_c100_custom_v1");'
           'localStorage.removeItem("wb_c100_bover_v1");'
           'localStorage.removeItem("wb_c100_groups_v1");}catch(e){}</script>')
    page = stripped[:j + 1] + pre + stripped[j + 1:]
    io.open(os.path.join(SITE, PAGE), 'w', encoding='utf-8').write(page)
    io.open(os.path.join(SITE, PROBE_PAGE), 'w', encoding='utf-8').write(
        page.replace('</body>', DBG.replace('@XLSX@', XLSX_NAME) + '</body>', 1))

    srv = serve()
    res = run_probe('http://localhost:%d/%s' % (PORT, PROBE_PAGE), timeout=260)
    print(res)
    print('=' * 78)
    srv.shutdown()
    shutil.rmtree(SITE, ignore_errors=True)
    print('（已清理探针临时目录 %s）' % os.path.relpath(SITE, HERE))
    bad = ('✗' in res) or ('JS错误: 无' not in res)
    print('结论：' + ('✗ 有用例失败' if bad else '✓ 全部通过（12 段）'))
    return 1 if bad else 0


if __name__ == '__main__':
    sys.exit(main())
