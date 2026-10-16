# -*- coding: utf-8 -*-
"""「我的商品上云」功能截图。走真实云端（localhost 在网关白名单内）。

注意：脚本会在云端创建探针数据，跑完需清理（见 main() 末尾打印的 id）。
"""
import functools
import http.server
import io
import os
import shutil
import socketserver
import subprocess
import sys
import threading
import time

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
SRC_PAGE = os.path.join(ROOT, 'deliver', '大宗商品日线对比图-2005-2026.html')
OUT = os.path.join(ROOT, 'deliver')
SITE = os.path.join(HERE, '_shot3_site')
XLSX = os.path.join(SITE, '_shot_wide.xlsx')
PORT = 8801
EDGE = '/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge'

HELPERS = r'''
<script>
(function(){
  function waitFor(fn, max, label){
    var n = 0;
    return new Promise(function(res, rej){
      (function tick(){
        var v = null;
        try{ v = fn(); }catch(e){ return rej(e); }
        if(v) return res(v);
        if(++n > max) return rej(new Error('timeout ' + label));
        requestAnimationFrame(tick);
      })();
    });
  }
  window.confirm = function(){ return true; };
  window.__T = {
    waitFor: waitFor,
    cloud: function(){ return waitFor(function(){ return CLOUD_STATE !== 'wait'; }, 40000, 'cloud'); },
    wideFile: function(){
      return fetch('_shot_wide.xlsx').then(function(r){ return r.arrayBuffer(); })
        .then(function(ab){ return new File([ab], '钼系品种.xlsx'); });
    }
  };
})();
</script>
'''

# 1) 首屏 + 工具栏云端徽标
S_MAIN = r'''
<script>
(async function(){
  try{ await __T.cloud(); }
  catch(e){ document.title = 'ERR ' + e.message; }
})();
</script>'''

# 2) 内置序列编辑面板（含「直接改价格数据」按钮）
S_BUI = r'''
<script>
(async function(){
  try{
    await __T.cloud();
    openEditBuiltin('gold');
    await __T.waitFor(function(){ return document.getElementById('eb-data'); }, 3000, 'panel');
  }catch(e){ document.title = 'ERR ' + e.message; }
})();
</script>'''

# 3) 「直接改价格数据」→ 覆盖编辑面板
S_ED = r'''
<script>
(async function(){
  try{
    await __T.cloud();
    openEditBuiltin('gold');
    await __T.waitFor(function(){ return document.getElementById('eb-data'); }, 3000, 'p1');
    document.getElementById('eb-data').click();
    await __T.waitFor(function(){ return document.getElementById('ed-save') && ED.code.indexOf('ov_') === 0; }, 4000, 'p2');
  }catch(e){ document.title = 'ERR ' + e.message; }
})();
</script>'''

# 4) 管理面板（自定义商品 + 内置覆盖）
S_MGR = r'''
<script>
(async function(){
  try{
    await __T.cloud();
    /* 传一条自定义商品 */
    var f = await __T.wideFile();
    openUpload();
    await __T.waitFor(function(){ return document.getElementById('fpick'); }, 3000, 'up');
    readUpFile(f);
    await __T.waitFor(function(){ return UP.list.length >= 1; }, 8000, 'parse');
    document.getElementById('up-ok').click();
    await __T.waitFor(function(){ return CUSTOM.length >= 1; }, 4000, 'commit');
    /* 再给内置 gold 造一条覆盖记录 */
    var g = BYCODE['gold'];
    if(!overlayOf('gold')){
      var b0 = g.pts0 || g.pts;
      CUSTOM.push({ c: 'ov_gold', n: g.n, m: '编辑内置数据', u: g.u || '', col: g.col || '#F5C542',
                    pts: b0.map(function(p){ return [p[0], p[1]]; }),
                    t0: b0[0][0], t1: b0[b0.length-1][0],
                    note: '编辑内置序列「' + g.n0 + '」的价格数据',
                    src: 'overlay', base: 'gold', t: Date.now() });
      afterCustomChange();
    }
    await __T.waitFor(function(){ return CUSTOM.length >= 2; }, 3000, 'mk');
    openMgr();
    await __T.waitFor(function(){ return document.querySelector('button[data-rm]'); }, 3000, 'mgr');
  }catch(e){ document.title = 'ERR ' + e.message; }
})();
</script>'''

SHOTS = [
    ('云端-工具栏状态.png',    's_main.html', S_MAIN, 1512, 1580, 45000),
    ('云端-内置编辑面板.png',  's_bui.html',  S_BUI,  1512, 1120, 45000),
    ('云端-直接改价格.png',    's_ed.html',   S_ED,   1512, 1420, 50000),
    ('云端-管理面板.png',      's_mgr.html',  S_MGR,  1512, 1000, 60000),
]


def make_xlsx():
    import datetime
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = '指标数据'
    ws.append(['指标名称', '现货价:钼精矿:Mo:45%:河南', '市场价:钼铁:FeMo60:上海市'])
    ws.append(['单位', '元/吨', '元/吨'])
    base = datetime.date(2023, 1, 18)
    i = 0
    while i < 200:
        d = base + datetime.timedelta(days=i)
        i += 1
        if d.weekday() >= 5:
            continue
        ws.append([d, round(5430 + i * 1.7, 1), round(341200 + i * 12.5, 1) if i < 130 else ''])
    for r in range(3, ws.max_row + 1):
        ws.cell(row=r, column=1).number_format = 'yyyy-mm-dd'
    wb.save(XLSX)
    return os.path.getsize(XLSX)


class Quiet(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a):
        pass


def serve():
    h = functools.partial(Quiet, directory=SITE)
    socketserver.TCPServer.allow_reuse_address = True
    httpd = socketserver.TCPServer(('127.0.0.1', PORT), h)
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    return httpd


def shoot(name, page, w, h, budget):
    out = os.path.join(OUT, name)
    if os.path.exists(out):
        os.remove(out)
    cmd = [EDGE, '--headless=new', '--no-sandbox', '--disable-gpu', '--disable-dev-shm-usage',
           '--hide-scrollbars', '--window-size=%d,%d' % (w, h),
           '--virtual-time-budget=%d' % budget, '--screenshot=' + out,
           'http://localhost:%d/%s' % (PORT, page)]
    subprocess.run(cmd, capture_output=True, timeout=300)
    return os.path.exists(out), (os.path.getsize(out) if os.path.exists(out) else 0)


def main():
    shutil.rmtree(SITE, ignore_errors=True)
    os.makedirs(SITE, exist_ok=True)
    n = make_xlsx()
    print('测试用 xlsx：%d 字节' % n)
    src = io.open(SRC_PAGE, encoding='utf-8').read()
    io.open(os.path.join(SITE, '_p.html'), 'w', encoding='utf-8').write(
        src.replace('</body>', HELPERS + '</body>', 1))
    for name, page, js, w, h, b in SHOTS:
        io.open(os.path.join(SITE, page), 'w', encoding='utf-8').write(
            src.replace('</body>', HELPERS + js + '</body>', 1))
    srv = serve()
    for name, page, js, w, h, b in SHOTS:
        t0 = time.time()
        ok, size = shoot(name, page, w, h, b)
        print('  %-24s %s  %8s  %4.1fs' % (name, '✓' if ok else '✗',
                                           ('%.0f KB' % (size / 1024)) if ok else '', time.time() - t0))
    srv.shutdown()
    shutil.rmtree(SITE, ignore_errors=True)
    print('（已清理临时目录）')
    print('\n⚠ 截图过程会向云端写入探针数据，跑完请清理 wb_c100_series 里的 u_shot_* / ov_gold')
    return 0


if __name__ == '__main__':
    sys.exit(main())
