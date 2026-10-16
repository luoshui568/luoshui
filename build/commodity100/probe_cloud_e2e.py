#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""「我的商品」云端入库 E2E 探针。

验证一条完整链路（浏览器跑在 localhost，走真实 CloudBase 网关）：
  0. 云端连通：WB_CFG 注入 → syncFromCloud() → CLOUD_STATE='on'
  1. Excel 宽表多列 → 一次性导入 → **落库 wb_c100_series**（Python 侧回读核对）
  2. 上传的数据可以直接改价格 → 保存 → 云端记录跟着变
  3. 内置序列「直接改价格」→ 生成 ov_<code> 覆盖记录 → 落库 → 图上数据被替换
  4. 删除自定义商品 → 云端软删 is_del=true
  5. 「恢复内置原始数据」→ 内置 pts 回到 pts0，云端覆盖记录被软删
  6. 第二段：重新打开页面 → 数据从云端恢复（不是靠 localStorage）

跑法：
  /Users/luoshui/.workbuddy/binaries/python/envs/default/bin/python probe_cloud_e2e.py
"""
import functools
import http.server
import io
import json
import os
import shutil
import socketserver
import subprocess
import sys
import threading
import time
import urllib.error
import urllib.request

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
SRC_PAGE = os.path.join(ROOT, 'deliver', '大宗商品日线对比图-2005-2026.html')
SITE = os.path.join(HERE, '_cloud_site')          # 独立临时目录：绝不碰 deploy_site/
PAGE = 'commodity-daily.html'
PROBE_PAGE = '_e2e_cloud.html'
PROBE_PAGE2 = '_e2e_cloud2.html'
XLSX_NAME = '_cloud_wide.xlsx'
XLSX = os.path.join(SITE, XLSX_NAME)
PORT = 8799
EDGE = '/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge'

CFG = json.load(io.open(os.path.join(ROOT, 'build', 'cloud.json'), encoding='utf-8'))
REST = 'https://%s.api.tcloudbasegateway.com/v1/rdb/rest/' % CFG['envId']


# ---------------------------------------------------------------- 云端直连（Python 侧核对）
def api(method, path, body=None):
    h = {'Authorization': 'Bearer ' + CFG['key']}
    if body is not None:
        h['Content-Type'] = 'application/json'
    r = urllib.request.Request(REST + path,
                               data=(json.dumps(body).encode() if body is not None else None),
                               headers=h, method=method)
    try:
        with urllib.request.urlopen(r, timeout=40) as resp:
            return resp.status, resp.read().decode('utf-8')
    except urllib.error.HTTPError as e:
        return e.code, e.read().decode('utf-8')


def cloud_rows():
    st, txt = api('GET', 'wb_c100_series?select=id,name,src,base,is_del,n,pts,up')
    if st != 200:
        print('  ! 读取云端失败', st, txt[:200])
        return []
    return json.loads(txt)


# ---------------------------------------------------------------- 造测试用 xlsx
def make_xlsx():
    import datetime
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = '指标数据'
    ws.append(['指标名称', '现货价:钼精矿:Mo:45%:河南', '市场价:钼铁:FeMo60:上海市',
               '市场价:钼铁:FeMo60:洛阳市:栾川', '备注'])
    ws.append(['单位', '元/吨', '元/吨', '元/吨', ''])
    base = datetime.date(2023, 1, 18)
    i = 0
    while i < 200:
        d = base + datetime.timedelta(days=i)
        i += 1
        if d.weekday() >= 5:
            continue
        ws.append([d,
                   round(5430 + i * 1.7, 1),
                   round(341200 + i * 12.5, 1) if i < 120 else '',
                   round(340700 + i * 11.1, 1) if i >= 70 else '',
                   '现货报价' if i % 7 == 0 else ''])
    for r in range(3, ws.max_row + 1):
        ws.cell(row=r, column=1).number_format = 'yyyy-mm-dd'
    wb.save(XLSX)
    return os.path.getsize(XLSX), ws.max_row


# ---------------------------------------------------------------- 探针脚本
DBG_BODY = r'''
<pre id="dbg" style="display:none"></pre>
<script>
(function(){
  var errs = [], out = [];
  window.addEventListener('error', function(e){ errs.push('ERR: ' + e.message); });
  window.addEventListener('unhandledrejection', function(e){
    errs.push('REJ: ' + ((e.reason && e.reason.message) || e.reason));
  });
  /* headless 下 confirm() 会阻塞，一律自动确认 */
  window.confirm = function(){ return true; };
  function log(k, v){ out.push(k + ' = ' + v); }
  /* 页面可能没有内置数据（build_opts.json → builtin:false）；内置相关用例据此跳过 */
  var HASB = (typeof HAS_BUILTIN === 'undefined') ? true : HAS_BUILTIN;
  function sec(t){ out.push(''); out.push('=== ' + t + ' ==='); }
  function txt(el){ return el ? (el.textContent || '').replace(/\s+/g, ' ').trim() : '(null)'; }
  function id(i){ return document.getElementById(i); }
  function isOpen(){ return !!id('mask') && id('mask').classList.contains('on'); }
  function title(){ var h = document.querySelector('#dlg h3'); return h ? h.textContent.replace(/\s+/g,' ').trim() : '(无标题)'; }
  function click(i){ var e = id(i); if(!e){ out.push('✗ 找不到按钮 ' + i); return false; } e.click(); return true; }
  function sleep(ms){ return new Promise(function(r){ setTimeout(r, ms); }); }
  function waitFor(fn, ms, tag){
    return new Promise(function(res, rej){
      var t0 = Date.now();
      (function poll(){
        var v; try{ v = fn(); }catch(e){ return rej(new Error(tag + ' 抛错: ' + e.message)); }
        if(v) return res(v);
        if(Date.now() - t0 > (ms || 3000)) return rej(new Error(tag + ' 超时(' + (ms || 3000) + 'ms)'));
        requestAnimationFrame(poll);
      })();
    });
  }
  function post(){
    fetch('/__cloud__', { method: 'POST',
      body: out.join('\n') + '\n\nJS错误: ' + (errs.length ? errs.join(' | ') : '无') });
  }

  (async function(){
    try{
      /* ---------------- 0. 云端配置与连通 ---------------- */
      sec('0. 云端配置与连通');
      log('WB_CFG.envId', (window.WB_CFG && window.WB_CFG.envId) || '(空)');
      log('WB_CFG.key 长度', ((window.WB_CFG && window.WB_CFG.key) || '').length);
      log('WBCloud.off()', WBCloud.off());
      await waitFor(function(){ return CLOUD_STATE !== 'wait'; }, 25000, '等云端同步');
      log('CLOUD_STATE', CLOUD_STATE);
      log('工具栏状态徽标', txt(id('cst')));
      var rows0 = await WBCloud.list();
      log('WBCloud.list() 可用', !!rows0);
      log('云端现有（未删）记录数', rows0 ? rows0.filter(function(r){ return !r.is_del; }).length : '—');

      /* ---------------- 1. 宽表多列 → 一次性导入 → 落库 ---------------- */
      sec('1. Excel 宽表多列 → 批量导入 → 落库');
      var n0 = CUSTOM.length;
      openUpload();
      await waitFor(function(){ return !!id('fpick'); }, 2000, '打开上传面板');
      var ab = await (await fetch('@XLSX@')).arrayBuffer();
      readUpFile(new File([ab], '钼系品种.xlsx'));
      await waitFor(function(){ return UP.list.length >= 3; }, 6000, '解析宽表');
      log('识别到的数值列', UP.list.length);
      log('默认勾选', UP.pick.size);
      click('up-ok');
      await waitFor(function(){ return CUSTOM.length >= n0 + 3; }, 4000, '导入落本机');
      var mine = CUSTOM.slice(n0);
      var mineIds = mine.map(function(s){ return s.c; });
      log('新导入条数', mine.length);
      log('新导入明细', mine.map(function(s){ return s.n + ':' + s.pts.length + '点'; }).join(' / '));
      await sleep(1500);                       // 等云端写入
      var rows1 = await WBCloud.list();
      var hit1 = mineIds.filter(function(x){ return rows1.some(function(r){ return r.id === x; }); });
      log('云端已存在这几条', hit1.length + '/' + mineIds.length);
      log('MINE_IDS', JSON.stringify(mineIds));

      /* ---------------- 2. 上传的数据可直接改价格 ---------------- */
      sec('2. 上传的数据 → 直接改价格 → 落库');
      var first = mine[0];
      openEdit(first.c);
      await waitFor(function(){ return isOpen() && !!id('ed-save'); }, 3000, '打开数据编辑');
      log('面板标题', title());
      var vinp = document.querySelector('#dlg input[data-f="v"]');
      var oldV = vinp.value, day0 = dayToStr(ED.pts[0][0]);
      vinp.value = '99999';
      vinp.dispatchEvent(new Event('change', { bubbles: true }));
      await waitFor(function(){ return ED.pts[0][1] === 99999; }, 1500, '改价未生效');
      log('改价', day0 + ' ' + oldV + ' → 99999');
      click('ed-save');
      await waitFor(function(){ return !isOpen(); }, 3000, '保存');
      await sleep(1500);
      var r2 = (await WBCloud.list()).filter(function(r){ return r.id === first.c; })[0];
      var r2v = r2 ? (r2.pts[0][1]) : null;
      log('云端回读该点价格', r2 ? (dayToStr(r2.pts[0][0]) + ' → ' + r2v) : '未找到');
      log('云端价格已更新', r2v === 99999);
      log('EDIT_ID', first.c);

      /* ---------------- 3. 内置序列直接改价格（覆盖记录） ---------------- */
      if(!HASB){
        sec('3. 内置序列 → 直接改价格');
        log('内置序列改价', '⊘ 本页无内置数据，跳过');
      }else{
      sec('3. 内置序列 → 直接改价格 → 生成覆盖记录 → 落库');
      var g0 = BYCODE['gold'].pts0[0][1];
      openEditBuiltin('gold');
      await waitFor(function(){ return isOpen() && !!id('eb-data'); }, 3000, '打开内置面板');
      log('内置面板标题', title());
      log('面板含「直接改价格数据」按钮', !!id('eb-data'));
      click('eb-data');
      await waitFor(function(){ return isOpen() && !!id('ed-save') && ED.code.indexOf('ov_') === 0; }, 3000, '进数据编辑');
      log('ED.code（覆盖记录 id）', ED.code);
      log('面板标题', title());
      var vinp2 = document.querySelector('#dlg input[data-f="v"]');
      vinp2.value = '88888';
      vinp2.dispatchEvent(new Event('change', { bubbles: true }));
      await waitFor(function(){ return ED.pts[0][1] === 88888; }, 1500, '改内置价未生效');
      click('ed-save');
      await waitFor(function(){ return !isOpen(); }, 3000, '保存覆盖');
      log('gold 覆盖记录挂上', !!BYCODE['gold'].ovId);
      log('gold 首点（现值 / 原始值）', BYCODE['gold'].pts[0][1] + ' / ' + BYCODE['gold'].pts0[0][1]);
      log('图上 gold 首点已替换', BYCODE['gold'].pts[0][1] === 88888 && g0 !== 88888);
      log('chip 出现「已改」标记', !!document.querySelector('.chip[data-c="gold"] .ovr'));
      var dsg = chart.data.datasets.filter(function(d){ return d._c === 'gold'; })[0];
      log('图表 gold 首点数据', dsg ? dsg.data[0].y : '未找到');
      await sleep(1500);
      var ovr = (await WBCloud.list()).filter(function(r){ return r.src === 'overlay' && r.base === 'gold'; })[0];
      log('云端覆盖记录', ovr ? (ovr.id + ' · ' + (ovr.pts ? ovr.pts.length : '?') + ' 点 · is_del=' + ovr.is_del) : '未找到');
      log('OV_ID', ovr ? ovr.id : '');
      }

      /* ---------------- 4. 删除自定义商品（软删） ---------------- */
      sec('4. 删除自定义商品 → 云端软删');
      var delId = mine[1].c;
      openMgr();
      await waitFor(function(){ return isOpen() && !!document.querySelector('button[data-rm="' + delId + '"]'); }, 3000, '打开管理面板');
      document.querySelector('button[data-rm="' + delId + '"]').click();
      await waitFor(function(){ return !CUSTOM.some(function(s){ return s.c === delId; }); }, 3000, '本机删除');
      log('本机已移除', !CUSTOM.some(function(s){ return s.c === delId; }));
      closeDlg();
      await sleep(1500);
      var d4 = (await WBCloud.list()).filter(function(r){ return r.id === delId; })[0];
      log('云端该条 is_del', d4 ? d4.is_del : '未找到');
      log('DEL_ID', delId);

      /* ---------------- 5. 恢复内置原始数据 ---------------- */
      if(!HASB){
        sec('5. 恢复内置原始数据');
        log('恢复内置原始数据', '⊘ 本页无内置数据，跳过');
      }else{
      sec('5. 恢复内置原始数据');
      var ovId = BYCODE['gold'].ovId;
      openEdit(ovId);
      await waitFor(function(){ return isOpen() && !!id('ed-del'); }, 3000, '打开覆盖编辑');
      log('面板标题', title());
      log('删除按钮文案', txt(id('ed-del')));
      click('ed-del');
      await waitFor(function(){ return /再点一次/.test(txt(id('ed-del'))); }, 2000, '二次确认');
      click('ed-del');
      await waitFor(function(){ return !isOpen(); }, 3000, '恢复完成');
      log('gold 首点已回到原始', BYCODE['gold'].pts[0][1] === BYCODE['gold'].pts0[0][1]);
      log('覆盖标记已清', !BYCODE['gold'].ovId);
      await sleep(1500);
      var d5 = (await WBCloud.list()).filter(function(r){ return r.id === ovId; })[0];
      log('云端覆盖记录 is_del', d5 ? d5.is_del : '未找到');
      }

      /* ---------------- 6. 剩余探针数据交给 Python 侧清理 ---------------- */
      sec('6. 备注');
      log('剩余自定义条数', CUSTOM.length);
      log('MINE_IDS_FINAL', JSON.stringify(CUSTOM.map(function(s){ return s.c; })));
    }catch(e){
      out.push('');
      out.push('★ 探针中断：' + (e && e.message ? e.message : e));
    }
    post();
  })();
})();
</script>
'''

DBG2 = r'''
<pre id="dbg" style="display:none"></pre>
<script>
(function(){
  var out = [], errs = [];
  window.addEventListener('error', function(e){ errs.push('ERR: ' + e.message); });
  function log(k, v){ out.push(k + ' = ' + v); }
  /* 页面可能没有内置数据（build_opts.json → builtin:false）；内置相关用例据此跳过 */
  var HASB = (typeof HAS_BUILTIN === 'undefined') ? true : HAS_BUILTIN;
  function sleep(ms){ return new Promise(function(r){ setTimeout(r, ms); }); }
  window.confirm = function(){ return true; };
  (async function(){
    try{
      out.push('=== 7. 重新打开页面 → 数据从云端恢复 ===');
      log('进页面时已清空本机缓存（清空后条数）', window.__LS0);
      log('清空动作确实发生过', !!window.__SEED);
      var t0 = Date.now();
      while(CLOUD_STATE === 'wait' && Date.now() - t0 < 25000){ await sleep(120); }
      log('CLOUD_STATE', CLOUD_STATE);
      log('CUSTOM 条数（只能来自云端）', CUSTOM.length);
      log('CUSTOM ids', JSON.stringify(CUSTOM.map(function(s){ return s.c; })));
      log('含内置覆盖记录', CUSTOM.filter(function(s){ return s.src === 'overlay'; }).length);
      log('页头点数', document.getElementById('ptn').textContent);
      log('状态徽标', (document.getElementById('cst') || {}).textContent || '');
    }catch(e){ out.push('★ 中断：' + (e && e.message ? e.message : e)); }
    fetch('/__cloud__', { method: 'POST',
      body: out.join('\n') + '\n\nJS错误: ' + (errs.length ? errs.join(' | ') : '无') });
  })();
})();
</script>
'''


# ---------------------------------------------------------------- 跑浏览器
RESULT = []


class Quiet(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def do_POST(self):
        n = int(self.headers.get('Content-Length') or 0)
        RESULT.append(self.rfile.read(n).decode('utf-8', 'ignore'))
        self.send_response(200)
        self.send_header('Content-Length', '2')
        self.end_headers()
        self.wfile.write(b'ok')


def serve():
    h = functools.partial(Quiet, directory=SITE)
    socketserver.TCPServer.allow_reuse_address = True
    httpd = socketserver.TCPServer(('127.0.0.1', PORT), h)
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    return httpd


def run_probe(url, timeout=180):
    del RESULT[:]
    cmd = [EDGE, '--headless=new', '--no-sandbox', '--disable-gpu', '--disable-dev-shm-usage',
           '--window-size=1512,1600', url]
    p = subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    t0 = time.time()
    while time.time() - t0 < timeout:
        if RESULT:
            break
        time.sleep(0.25)
    time.sleep(0.4)
    try:
        p.terminate(); p.wait(timeout=8)
    except Exception:
        try: p.kill()
        except Exception: pass
    return RESULT[0] if RESULT else '（%.0f 秒内没收到探针回传）' % timeout


def main():
    shutil.rmtree(SITE, ignore_errors=True)
    os.makedirs(SITE, exist_ok=True)

    print('=' * 76)
    print('「我的商品」云端入库 E2E')
    print('=' * 76)
    before = {r['id'] for r in cloud_rows()}
    print('跑之前云端记录数:', len(before))

    size, rows = make_xlsx()
    print('测试用 xlsx：%s（%d 字节 · %d 行）' % (XLSX_NAME, size, rows))

    src = io.open(SRC_PAGE, encoding='utf-8').read()
    assert 'window.WB_CFG = { envId: "fdels568' in src, '构建产物里没有注入云端配置'
    io.open(os.path.join(SITE, PAGE), 'w', encoding='utf-8').write(src)
    io.open(os.path.join(SITE, PROBE_PAGE), 'w', encoding='utf-8').write(
        src.replace('</body>', DBG_BODY.replace('@XLSX@', XLSX_NAME) + '</body>', 1))

    # 第二段：刻意把 localStorage 里的自定义商品清空，证明数据确实来自云端
    i = src.find('<body')
    j = src.find('>', i)
    pre = ('<script>try{localStorage.removeItem("wb_c100_custom_v1");'
           'window.__LS0=(JSON.parse(localStorage.getItem("wb_c100_custom_v1")||"[]")).length;'
           'window.__SEED=1;}catch(e){window.__LS0=-1;}</script>')
    page2 = src[:j + 1] + pre + src[j + 1:]
    io.open(os.path.join(SITE, PROBE_PAGE2), 'w', encoding='utf-8').write(
        page2.replace('</body>', DBG2 + '</body>', 1))

    srv = serve()
    print('\n' + '=' * 76)
    seg1 = run_probe('http://localhost:%d/%s' % (PORT, PROBE_PAGE), timeout=200)
    print(seg1)
    print('=' * 76)
    seg2 = run_probe('http://localhost:%d/%s' % (PORT, PROBE_PAGE2), timeout=120)
    print(seg2)
    print('=' * 76)
    srv.shutdown()

    # ---------------- Python 侧独立回读核对 ----------------
    after = cloud_rows()
    new = [r for r in after if r['id'] not in before]
    print('\n【Python 侧独立核对】清理前云端新增记录 %d 条：' % len(new))
    for r in sorted(new, key=lambda x: x['id']):
        print('   %-22s %-8s base=%-6s is_del=%-5s %s点  %s'
              % (r['id'], r['src'], r['base'] or '-', r['is_del'], r['n'], r['name'][:20]))

    # 收尾：把探针造的数据软删掉（硬删由 MCP 侧执行）
    ids = [r['id'] for r in new]
    if ids:
        q = ','.join('"%s"' % x for x in ids)
        st, _ = api('PATCH', 'wb_c100_series?id=in.(%s)' % q, {'is_del': True})
        print('\n已把 %d 条探针记录标记为软删（HTTP %s）；id 列表：' % (len(ids), st))
        print(','.join(ids))
    else:
        print('\n（没有新增记录 —— 可能整条链路没跑通）')

    shutil.rmtree(SITE, ignore_errors=True)
    print('\n（已清理探针临时目录 %s）' % os.path.relpath(SITE, HERE))
    return 0


if __name__ == '__main__':
    sys.exit(main())
