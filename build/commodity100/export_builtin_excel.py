#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把内置 31 条日线序列（价格 + 时间）导出成一张「宽表 Excel」，并可选地
按「上传形式」写进云端表 wb_c100_series。

宽表形态刻意与页面上传面板的识别规则对齐（见 template_daily.html 的 guessCols）：
    第 1 行   品种名（'日期' 打头，其余 31 列各一个品种）
    第 2 行   单位（'单位' 打头）
    第 3 行起 第一列日期（真日期型，number_format yyyy-mm-dd），其余列是价格
所以导出的文件可以直接拖进「⬆ 上传 Excel」，会被识别成 31 列、一次全部导入。

内置数据是打包在页面里的紧凑格式：
    data/daily_data.json → groups[].series[].d = [[dayNo, priceCents], ...]
    dayNo      = 距 1970-01-01 的天数
    priceCents = 价格 × 100 取整

用法：
    python export_builtin_excel.py                 # 只导出到桌面
    python export_builtin_excel.py --out X.xlsx    # 指定输出路径
    python export_builtin_excel.py --push          # 导出 + 入库为「上传形式」
"""
import argparse
import datetime as dt
import json
import os
import sys
import time
import urllib.request

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, 'data', 'daily_data.json')
CLOUD_JSON = os.path.abspath(os.path.join(HERE, '..', 'cloud.json'))
TABLE = 'wb_c100_series'
EPOCH = dt.date(1970, 1, 1)


def day_to_date(n):
    return EPOCH + dt.timedelta(days=int(n))


def load_series():
    d = json.load(open(DATA, encoding='utf-8'))
    out = []
    for g in d['groups']:
        for s in g['series']:
            pts = [[int(p[0]), int(p[1]) / 100.0] for p in s['d']]
            pts.sort(key=lambda x: x[0])
            out.append({
                'code': s['c'], 'name': s['n'], 'unit': s['u'], 'market': s['m'],
                'color': s['col'], 'breed': s['b'], 'origin': s.get('src', ''),
                'symbol': s.get('sym', ''), 'group': g['name'], 'pts': pts,
            })
    return out, d.get('built', '')


def build_grid(series):
    """把 N 条序列摊成宽表：第一列日期（并集），其余各列按品种铺开，缺的留空。"""
    days = set()
    for s in series:
        days.update(p[0] for p in s['pts'])
    days = sorted(days)
    idx = {d: i for i, d in enumerate(days)}
    cols = []
    for s in series:
        col = [None] * len(days)
        for d, v in s['pts']:
            col[idx[d]] = v
        cols.append(col)
    return days, cols


def write_xlsx(path, series, days, cols, built):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = '日线数据'

    hdr = Font(bold=True, color='FFFFFF')
    fill = PatternFill('solid', fgColor='C8102E')
    sub = PatternFill('solid', fgColor='F2E3E6')

    ws.append(['日期'] + [s['name'] for s in series])
    ws.append(['单位'] + [s['unit'] for s in series])
    for c in range(1, len(series) + 2):
        ws.cell(row=1, column=c).font = hdr
        ws.cell(row=1, column=c).fill = fill
        ws.cell(row=2, column=c).fill = sub
        ws.cell(row=2, column=c).alignment = Alignment(horizontal='center')

    for i, dn in enumerate(days):
        row = [day_to_date(dn)]
        for col in cols:
            v = col[i]
            row.append(None if v is None else v)
        ws.append(row)
    for r in range(3, ws.max_row + 1):
        ws.cell(row=r, column=1).number_format = 'yyyy-mm-dd'
        ws.cell(row=r, column=1).alignment = Alignment(horizontal='left')

    ws.freeze_panes = 'B3'
    ws.column_dimensions['A'].width = 12
    for c in range(2, len(series) + 2):
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 15

    # 第二张表：品种清单（数据表必须排第一，上传面板默认取第一个工作表）
    ws2 = wb.create_sheet('品种说明')
    ws2.append(['#', '品种', '内置代码', '分类', '市场', '单位', '颜色',
                '起始日', '结束日', '数据点', '数据来源'])
    for c in range(1, 12):
        ws2.cell(row=1, column=c).font = hdr
        ws2.cell(row=1, column=c).fill = fill
    for i, s in enumerate(series, 1):
        p = s['pts']
        ws2.append([i, s['name'], s['code'], s['group'], s['market'], s['unit'], s['color'],
                    day_to_date(p[0][0]).isoformat() if p else '',
                    day_to_date(p[-1][0]).isoformat() if p else '',
                    len(p), s['origin']])
    for c, w in zip('ABCDEFGHIJK', (4, 18, 12, 10, 14, 12, 10, 12, 12, 10, 22)):
        ws2.column_dimensions[c].width = w
    ws2.freeze_panes = 'A2'

    wb.save(path)
    return os.path.getsize(path), ws.max_row


# ------------------------------------------------------------------ 入库
def rest_base():
    cfg = json.load(open(CLOUD_JSON, encoding='utf-8'))
    env = cfg['envId']
    return 'https://%s.api.tcloudbasegateway.com/v1/rdb/rest/' % env, cfg['key']


def post(rows, key, base):
    req = urllib.request.Request(
        base + TABLE + '?on_conflict=id',
        data=json.dumps(rows, ensure_ascii=False).encode('utf-8'),
        headers={
            'Authorization': 'Bearer ' + key,
            'Content-Type': 'application/json',
            'Prefer': 'resolution=merge-duplicates,return=minimal',
        },
        method='POST')
    with urllib.request.urlopen(req, timeout=180) as r:
        return r.status


def to_rows(series, stamp, owner):
    rows = []
    for i, s in enumerate(series):
        p = s['pts']
        rows.append({
            'id': 'u_%s%x' % (stamp, i),
            'name': s['name'], 'market': '自定义上传', 'unit': s['unit'], 'color': s['color'],
            't0': p[0][0], 't1': p[-1][0], 'n': len(p), 'pts': p,
            'note': '内置序列「%s」导出为上传形式（原内置 code=%s）' % (s['name'], s['code']),
            'src': 'upload', 'base': '', 'owner': owner, 'is_del': False,
            'up': int(time.time() * 1000),
        })
    return rows


def get(path, key, base):
    req = urllib.request.Request(base + path, headers={'Authorization': 'Bearer ' + key})
    return json.load(urllib.request.urlopen(req, timeout=60))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--out', default=os.path.expanduser('~/Desktop/大宗商品日线-31品种-宽表.xlsx'))
    ap.add_argument('--push', action='store_true', help='同时按「上传形式」写入云端 wb_c100_series')
    ap.add_argument('--owner', default='超级管理员')
    a = ap.parse_args()

    series, built = load_series()
    days, cols = build_grid(series)
    total = sum(len(s['pts']) for s in series)
    print('内置序列 %d 条 / 数据点 %d 个 / 日期并集 %d 天（%s → %s）'
          % (len(series), total, len(days),
             day_to_date(days[0]).isoformat(), day_to_date(days[-1]).isoformat()))

    size, rows = write_xlsx(a.out, series, days, cols, built)
    print('✓ 已导出：%s' % a.out)
    print('  %.1f KB · %d 行 × %d 列（1 列日期 + %d 个品种）'
          % (size / 1024.0, rows, len(series) + 1, len(series)))

    if not a.push:
        print('\n（未加 --push，只导出文件。要同时入库为「上传形式」就加 --push）')
        return 0

    base, key = rest_base()
    stamp = base36(int(time.time() * 1000))
    rowsj = to_rows(series, stamp, a.owner)
    st = post(rowsj, key, base)
    print('\n✓ 入库 %d 条（HTTP %s）' % (len(rowsj), st))
    ids = [r['id'] for r in rowsj]
    meta = {
        'stamp': stamp, 'owner': a.owner, 'xlsx': a.out, 'built': built,
        'total_points': total, 'days': len(days),
        'series': [{'id': r['id'], 'name': r['name'], 'unit': r['unit'], 'color': r['color'],
                    'n': r['n'], 't0': r['t0'], 't1': r['t1']} for r in rowsj],
    }
    mpath = os.path.join(HERE, '_export_meta.json')
    json.dump(meta, open(mpath, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print('  清单：%s' % mpath)
    got = get(TABLE + '?select=id,name,n,t0,t1,src,is_del&id=in.(%s)&order=name' % ','.join(ids), key, base)
    print('  回读核对：%d/%d 条在位' % (len(got), len(ids)))
    bad = [r for r in got if r['is_del'] or not r['n']]
    print('  异常记录：%d 条' % len(bad))
    print('  点合计：%d' % sum(r['n'] for r in got))
    print('  新增 id 段：%s ... %s' % (ids[0], ids[-1]))
    return 0


def base36(n):
    alpha = '0123456789abcdefghijklmnopqrstuvwxyz'
    s = ''
    while n:
        n, r = divmod(n, 36)
        s = alpha[r] + s
    return s or '0'


if __name__ == '__main__':
    sys.exit(main())
