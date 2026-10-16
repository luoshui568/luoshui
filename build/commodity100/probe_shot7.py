# -*- coding: utf-8 -*-
"""「分组管理（增删改查）+ 导入面板分组列」功能截图。

★ 截图站点把 window.WB_CFG 掏空（本机模式）：演示用的分组只写进本页 localStorage，
  一条都不会落到 wb_c100_groups；商品数据是**从云端只读拉下来的真实商品**，
  所以截图里的名字和用户看到的完全一致。
"""
import datetime
import functools
import http.server
import io
import json
import os
import re
import shutil
import socketserver
import subprocess
import sys
import threading
import time
import urllib.request

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
SRC_PAGE = os.path.join(ROOT, 'deliver', '大宗商品日线对比图-2005-2026.html')
OUT = os.path.join(ROOT, 'deliver')
SITE = os.path.join(HERE, '_shot7_site')
CLOUD_JSON = os.path.abspath(os.path.join(HERE, '..', 'cloud.json'))
XLSX_NAME = 'demo_wide.xlsx'
PORT = 8813
EDGE = '/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge'


def fetch_cloud(limit_rows=8):
    cfg = json.load(io.open(CLOUD_JSON, encoding='utf-8'))
    rest = 'https://%s.api.tcloudbasegateway.com/v1/rdb/rest/' % cfg['envId']
    url = rest + 'wb_c100_series?select=id,name,unit,color,tag,note,src,base,pts,up&is_del=eq.false&order=up.asc'
    r = urllib.request.Request(url, headers={'Authorization': 'Bearer ' + cfg['key']})
    rows = json.load(urllib.request.urlopen(r, timeout=60))
    recs = []
    for x in rows[:limit_rows]:
        pts = x.get('pts') or []
        if len(pts) < 2:
            continue
        recs.append({'c': x['id'], 'n': x['name'], 'u': x.get('unit') or '',
                     'col': x.get('color') or '#FF4D6D', 'tag': x.get('tag') or '',
                     'note': x.get('note') or '', 'src': 'upload',
                     'pts': [[int(p[0]), p[1]] for p in pts], 't': 1})
    return recs


def make_xlsx():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = '日线'
    ws.append(['日期', '黄金 · COMEX', '白银 · COMEX', '钼精矿 · Mo · 45%', '稀土 · 氧化镨钕'])
    ws.append(['单位', '美元/盎司', '美元/盎司', '元/吨', '元/吨'])
    d = datetime.date(2026, 1, 1)
    i = 0
    while i < 200:
        if d.weekday() < 5:
            ws.append([d, round(4000 + i * 1.7, 2), round(48 + i * 0.09, 3),
                       round(3900 + i * 2.1, 1), round(520000 + i * 260.0, 1)])
        d += datetime.timedelta(days=1)
        i += 1
    for r in range(3, ws.max_row + 1):
        ws.cell(row=r, column=1).number_format = 'yyyy-mm-dd'
    wb.save(os.path.join(SITE, XLSX_NAME))


def seed_script(recs, groups=None):
    js = ('try{localStorage.setItem("wb_c100_custom_v1",'
          + json.dumps(json.dumps(recs, ensure_ascii=False)) + ');')
    if groups:
        js += ('localStorage.setItem("wb_c100_groups_v1",'
               + json.dumps(json.dumps(groups, ensure_ascii=False)) + ');')
    js += '}catch(e){}'
    return '<script>' + js + '</script>'


def demo_group(name, cid):
    """演示用的已登记分组（与本机模式一致：只活在 localStorage 里）"""
    return {'c': cid, 'n': name, 'sort': 99, 'color': '', 'note': '',
            'owner': '演示', 't': 1}


COMMON = '''
<script>
function sleep(ms){ return new Promise(function(r){ setTimeout(r, ms); }); }
function waitFor(fn, ms, tag){
  return new Promise(function(res, rej){
    var t0 = Date.now();
    (function poll(){
      var v; try{ v = fn(); }catch(e){ return rej(new Error(tag + ' 抛错: ' + e.message)); }
      if(v) return res(v);
      if(Date.now() - t0 > (ms || 3000)) return rej(new Error(tag + ' 超时'));
      requestAnimationFrame(poll);
    })();
  });
}
function setVal(el, v){
  el.value = v;
  el.dispatchEvent(new Event('input',  { bubbles: true }));
  el.dispatchEvent(new Event('change', { bubbles: true }));
}
window.confirm = function(){ return true; };
</script>'''

# ① 分组管理面板（顺带演示「新建空分组」）
S_GRP = COMMON + '''
<script>
(async function(){
  try{
    openMgr();
    await sleep(400);
    var b = document.getElementById('mg-grp');
    if(!b){ document.title = 'ERR 没有分组管理入口'; return; }
    b.click();
    await sleep(400);
    var n = document.getElementById('gm-new');
    if(!n){ document.title = 'ERR 分组面板没出来'; return; }
    setVal(n, '稀土');
    document.getElementById('gm-add').click();
    await sleep(400);
    if(!document.getElementById('gm-add')) document.title = 'ERR 新建后又没了';
  }catch(e){ document.title = 'ERR ' + e.message; }
})();
</script>'''

# ② 管理面板（分组管理入口）
S_MGR = COMMON + '''
<script>
(async function(){
  try{ openMgr(); await sleep(500); }catch(e){ document.title = 'ERR ' + e.message; }
})();
</script>'''

# ③ 上传面板：每一行单独的「分组」列
S_UP = COMMON + '''
<script>
(async function(){
  try{
    document.getElementById('btn-up').click();
    await waitFor(function(){ return !!document.getElementById('fpick'); }, 5000, '上传面板');
    var ab = await (await fetch('@XLSX@')).arrayBuffer();
    readUpFile(new File([ab], '商品日线数据.xlsx'));
    await waitFor(function(){ return UP.list && UP.list.length >= 4; }, 25000, '解析 xlsx');
    var cs = UP.list.map(function(it){ return it.c; });
    function setTag(i, t){
      var s = document.querySelector('#up-list select[data-tg="' + cs[i] + '"]');
      if(!s) return;
      s.value = t;
      s.dispatchEvent(new Event('change', { bubbles: true }));
    }
    setTag(0, '贵金属');
    setTag(1, '贵金属');
    setTag(2, '小金属');
    setTag(3, '稀土');          /* 稀土是刚在分组管理里新建的空分组（本页也从缓存里读到了） */
  }catch(e){ document.title = 'ERR ' + e.message; }
})();
</script>'''


class Quiet(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a):
        pass


def serve():
    h = functools.partial(Quiet, directory=SITE)
    socketserver.TCPServer.allow_reuse_address = True
    httpd = socketserver.TCPServer(('127.0.0.1', PORT), h)
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    return httpd


def shoot(name, page, w, h, budget):
    out = os.path.join(OUT, name)
    if os.path.exists(out):
        os.remove(out)
    cmd = [EDGE, '--headless=new', '--no-sandbox', '--disable-gpu', '--disable-dev-shm-usage',
           '--hide-scrollbars', '--window-size=%d,%d' % (w, h),
           '--virtual-time-budget=%d' % budget, '--screenshot=' + out,
           'http://localhost:%d/%s' % (PORT, page)]
    r = subprocess.run(cmd, capture_output=True, timeout=300)
    err = (r.stderr or b'').decode('utf-8', 'ignore')
    if '<title>ERR' in err:
        pass
    return os.path.exists(out), (os.path.getsize(out) if os.path.exists(out) else 0)


def main():
    recs = fetch_cloud()
    print('演示商品（只读自云端）：%d 条 —— %s' % (len(recs), ' / '.join(x['n'] for x in recs)))
    print('  标签分布：%s' % (' / '.join(sorted(set((x['tag'] or '（空）') for x in recs)))))
    shutil.rmtree(SITE, ignore_errors=True)
    os.makedirs(SITE, exist_ok=True)
    make_xlsx()
    src = io.open(SRC_PAGE, encoding='utf-8').read()
    stripped, cnt = re.subn(r'window\.WB_CFG\s*=\s*\{[^}]*\};',
                            'window.WB_CFG = null;   /* 截图站点：本机模式 */', src)
    assert cnt == 1, '未能掏空 WB_CFG（匹配 %d 次）' % cnt
    i = stripped.find('<body')
    j = stripped.find('>', i)
    head, body = stripped[:j + 1], stripped[j + 1:]
    seed = seed_script(recs)
    # 导入面板那张：让「稀土」这个刚建的空分组也在这页可见（本机模式，只进 localStorage）
    seed_up = seed_script(recs, [demo_group('稀土', 'g_demo_xt')])

    io.open(os.path.join(SITE, 's_grp.html'), 'w', encoding='utf-8').write(
        head + seed + body.replace('</body>', S_GRP + '</body>', 1))
    io.open(os.path.join(SITE, 's_mgr.html'), 'w', encoding='utf-8').write(
        head + seed + body.replace('</body>', S_MGR + '</body>', 1))
    io.open(os.path.join(SITE, 's_up.html'), 'w', encoding='utf-8').write(
        head + seed_up + body.replace('</body>', S_UP.replace('@XLSX@', XLSX_NAME) + '</body>', 1))

    SHOTS = [
        ('分组管理-增删改查.png',   's_grp.html', 1512, 1080, 45000),
        ('分组管理-管理面板入口.png', 's_mgr.html', 1512,  980, 45000),
        ('导入-每行可选分组.png',    's_up.html',  1512, 1440, 60000),
    ]
    srv = serve()
    bad = 0
    for name, page, w, h, b in SHOTS:
        t0 = time.time()
        ok, size = shoot(name, page, w, h, b)
        bad += 0 if ok else 1
        print('  %-30s %s  %8s  %4.1fs' % (name, '✓' if ok else '✗',
                                           ('%.0f KB' % (size / 1024)) if ok else '', time.time() - t0))
    srv.shutdown()
    shutil.rmtree(SITE, ignore_errors=True)
    print('（已清理临时目录；本机模式截图，未写入 wb_c100_groups / wb_c100_series）')
    return 1 if bad else 0


if __name__ == '__main__':
    sys.exit(main())
