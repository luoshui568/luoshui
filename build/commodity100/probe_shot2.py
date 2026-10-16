#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""本次三条改动出视觉截图：多列批量导入 / 内置序列编辑 / 工具栏按钮。

和 probe_v2_e2e.py 一样：临时目录自建、不碰 deploy_site/。
异步等待一律用 requestAnimationFrame（虚拟时间会把 setTimeout 瞬间跳过）。
"""
import functools
import http.server
import io
import os
import shutil
import socketserver
import subprocess
import sys
import threading

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
SRC_PAGE = os.path.join(ROOT, 'deliver', '大宗商品日线对比图-2005-2026.html')
OUT = os.path.join(ROOT, 'deliver')
SITE = os.path.join(HERE, '_shot2_site')
XLSX = os.path.join(SITE, '_shot_wide.xlsx')
PORT = 8796
EDGE = '/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge'

HELPERS = r'''
<script>
(function(){
  function waitFor(fn, max, label){
    var n = 0;
    return new Promise(function(res, rej){
      (function tick(){
        var v = null;
        try{ v = fn(); }catch(e){ return rej(e); }
        if(v) return res(v);
        if(++n > max) return rej(new Error('timeout ' + label));
        requestAnimationFrame(tick);
      })();
    });
  }
  function wideFile(){
    return fetch('_shot_wide.xlsx').then(function(r){ return r.arrayBuffer(); })
      .then(function(ab){ return new File([ab], '钼系品种.xlsx'); });
  }
  window.__T = { waitFor: waitFor, wideFile: wideFile };
})();
</script>
'''

S_UP = r'''
<script>
(async function(){
  try{
    var f = await __T.wideFile();
    openUpload();
    readUpFile(f);
    await __T.waitFor(function(){ return UP.list && UP.list.length; }, 900, 'parse');
    var pv = document.getElementById('up-prev'); if(pv) pv.style.display = 'none';
  }catch(e){ document.title = 'ERR ' + e.message; }
})();
</script>'''

S_BUILTIN = r'''
<script>
(async function(){
  try{
    document.querySelector('.chip[data-c="gold"] .ed').click();
    await __T.waitFor(function(){ return !!document.getElementById('eb-name'); }, 300, 'eb');
    var box = document.querySelector('#dlg .dwarn');
    if(box) box.style.display = 'none';
  }catch(e){ document.title = 'ERR ' + e.message; }
})();
</script>'''

S_BAR = r'''
<script>
(async function(){
  try{
    /* 顶到页首：chips 分组与工具栏就在页头下方，高清视口一次拍全 */
    window.scrollTo(0, 0);
    document.getElementById('chips').style.outline = 'none';
  }catch(e){ document.title = 'ERR ' + e.message; }
})();
</script>'''

SHOTS = [
    ('v2-多列批量导入.png', 'shot_up.html', S_UP, 1512, 1320),
    ('v2-内置序列编辑.png', 'shot_eb.html', S_BUILTIN, 1512, 1180),
    ('v2-选择区与工具栏.png', 'shot_bar.html', S_BAR, 1512, 1010),
]


def make_xlsx():
    import datetime
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = '指标数据'
    ws.append(['指标名称', '现货价:钼精矿:Mo:45%:河南', '市场价:钼铁:FeMo60:上海市',
               '市场价:钼铁:FeMo60:洛阳市:栾川', '备注'])
    ws.append(['单位', '元/吨', '元/吨', '元/吨', ''])
    base = datetime.date(2023, 1, 18)
    i = 0
    while i < 260:
        d = base + datetime.timedelta(days=i)
        i += 1
        if d.weekday() >= 5:
            continue
        ws.append([d,
                   round(5430 + i * 1.7, 1),
                   round(341200 + i * 12.5, 1) if i < 160 else '',
                   round(340700 + i * 11.1, 1) if i >= 90 else '',
                   '现货报价' if i % 7 == 0 else ''])
    for r in range(3, ws.max_row + 1):
        ws.cell(row=r, column=1).number_format = 'yyyy-mm-dd'
    wb.save(XLSX)
    return os.path.getsize(XLSX)


class Quiet(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a):
        pass


def serve():
    h = functools.partial(Quiet, directory=SITE)
    socketserver.TCPServer.allow_reuse_address = True
    httpd = socketserver.TCPServer(('127.0.0.1', PORT), h)
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    return httpd


def shoot(name, page, w, h):
    out = os.path.join(OUT, name)
    cmd = [EDGE, '--headless=new', '--no-sandbox', '--disable-gpu', '--disable-dev-shm-usage',
           '--hide-scrollbars', '--window-size=%d,%d' % (w, h),
           '--virtual-time-budget=22000', '--screenshot=' + out,
           'http://localhost:%d/%s' % (PORT, page)]
    subprocess.run(cmd, capture_output=True, timeout=300)
    return os.path.exists(out), (os.path.getsize(out) if os.path.exists(out) else 0)


def main():
    shutil.rmtree(SITE, ignore_errors=True)
    os.makedirs(SITE, exist_ok=True)
    n = make_xlsx()
    print('测试用 xlsx：%d 字节' % n)
    src = io.open(SRC_PAGE, encoding='utf-8').read()
    io.open(os.path.join(SITE, '_p.html'), 'w', encoding='utf-8').write(
        src.replace('</body>', HELPERS + '</body>', 1))
    for name, page, js, w, h in SHOTS:
        io.open(os.path.join(SITE, page), 'w', encoding='utf-8').write(
            src.replace('</body>', HELPERS + js + '</body>', 1))
    srv = serve()
    for name, page, js, w, h in SHOTS:
        ok, size = shoot(name, page, w, h)
        print('  %-22s %s  %s' % (name, '✓' if ok else '✗', ('%.0f KB' % (size / 1024)) if ok else ''))
    srv.shutdown()
    shutil.rmtree(SITE, ignore_errors=True)
    print('（已清理临时目录）')
    return 0


if __name__ == '__main__':
    sys.exit(main())
