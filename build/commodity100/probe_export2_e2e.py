#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""「页面无内置数据 + 备份面板导出 Excel/CSV」E2E 探针。

★ 本探针跑在**本机模式**：main() 把注入的 window.WB_CFG 掏空 → WBCloud.off() 为 true，
  所有写入退回本机，wb_c100_series 一条都不动。

覆盖：
  0. 页面形态：HAS_BUILTIN=false、内置分组 0 个、来源说明隐藏、空图提示可见、页头无「内置」字样
  1. 导入 3 列 → 自动全选、图表有 dataset、空图提示消失
  2. 备份面板：Excel / CSV / json 三个按钮齐全，点下去不抛错（下载被拦下来统计）
  3. 导出的 xlsx：结构自检（工作表/表头/日期/列数）
  4. 导出的 xlsx **再拖回上传面板** → 列名/单位/点数与源逐列一致（往返闭合）
  5. CSV 文本形状：BOM、行数、表头、日期列
  6. 把 xlsx 字节回传 Python，用 openpyxl 独立打开再核一遍（不走页面自己的解析器）
  7. 清空后空图提示回来
  8. 全程零 JS 错误

跑法：
  /Users/luoshui/.workbuddy/binaries/python/envs/default/bin/python probe_export2_e2e.py
"""
import base64
import functools
import http.server
import io
import json
import os
import re
import shutil
import socketserver
import subprocess
import sys
import threading
import time

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
SRC_PAGE = os.path.join(ROOT, 'deliver', '大宗商品日线对比图-2005-2026.html')
SITE = os.path.join(HERE, '_export2_site')        # 独立临时目录：绝不碰 deploy_site/
PAGE = 'commodity-daily.html'
PROBE_PAGE = '_e2e_export2.html'
XLSX_NAME = '_export2_src.xlsx'
XLSX = os.path.join(SITE, XLSX_NAME)
OUT_XLSX = '/tmp/_probe_export2.xlsx'
PORT = 8805
EDGE = '/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge'
RESULT = []
XLSXB64 = {}


def make_xlsx():
    """源数据：一列日期 + 3 个数值列（名称/单位各不相同，便于核对表头与单位行）"""
    import datetime
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = '日线'
    ws.append(['日期', '黄金 · COMEX', '白银 · COMEX', '钼铁 · FeMo60', '备注'])
    ws.append(['单位', '美元/盎司', '美元/盎司', '元/吨', ''])
    d = datetime.date(2024, 1, 1)
    i = 0
    while i < 90:
        if d.weekday() < 5:
            ws.append([d, round(2000 + i * 3.3, 2), round(24 + i * 0.11, 3),
                       round(138000 + i * 220.0, 1), '现货' if i % 11 == 0 else ''])
        d += datetime.timedelta(days=1)
        i += 1
    for r in range(3, ws.max_row + 1):
        ws.cell(row=r, column=1).number_format = 'yyyy-mm-dd'
    wb.save(XLSX)
    return os.path.getsize(XLSX), ws.max_row


DBG = r'''
<pre id="dbg" style="display:none"></pre>
<script>
(function(){
  var errs = [], out = [];
  window.addEventListener('error', function(e){ errs.push('ERR: ' + e.message); });
  window.addEventListener('unhandledrejection', function(e){
    errs.push('REJ: ' + ((e.reason && e.reason.message) || e.reason));
  });
  window.confirm = function(){ return true; };
  window.prompt = function(){ return ''; };

  function log(k, v){ out.push(k + ' = ' + v); }
  function sec(t){ out.push(''); out.push('=== ' + t + ' ==='); }
  function id(i){ return document.getElementById(i); }
  function txt(el){ return el ? (el.textContent || '').replace(/\s+/g, ' ').trim() : '(null)'; }
  function click(el){ if(!el){ out.push('✗ 找不到元素'); return false; } el.click(); return true; }
  function clickId(i){ return click(id(i)); }
  function sleep(ms){ return new Promise(function(r){ setTimeout(r, ms); }); }
  function waitFor(fn, ms, tag){
    return new Promise(function(res, rej){
      var t0 = Date.now();
      (function poll(){
        var v; try{ v = fn(); }catch(e){ return rej(new Error(tag + ' 抛错: ' + e.message)); }
        if(v) return res(v);
        if(Date.now() - t0 > (ms || 3000)) return rej(new Error(tag + ' 超时(' + (ms || 3000) + 'ms)'));
        requestAnimationFrame(poll);
      })();
    });
  }
  function setVal(el, v){
    el.value = v;
    el.dispatchEvent(new Event('input',  { bubbles: true }));
    el.dispatchEvent(new Event('change', { bubbles: true }));
  }
  function vis(el){ return !!el && getComputedStyle(el).display !== 'none'; }
  function b64(u8){
    var s = '', CH = 0x8000;
    for(var i = 0; i < u8.length; i += CH) s += String.fromCharCode.apply(null, u8.subarray(i, i + CH));
    return btoa(s);
  }
  async function importXlsx(tag){
    clickId('btn-up');
    await waitFor(function(){ return !!id('fpick'); }, 4000, '打开上传面板');
    var ab = await (await fetch('@XLSX@')).arrayBuffer();
    readUpFile(new File([ab], '导出测试源.xlsx'));
    await waitFor(function(){ return UP.list && UP.list.length >= 3; }, 20000, '解析 xlsx');
    var n0 = CUSTOM.length;
    if(tag !== null && id('up-tag')) setVal(id('up-tag'), tag);
    clickId('up-ok');
    await waitFor(function(){ return CUSTOM.length > n0; }, 6000, '导入落地');
    return CUSTOM.length - n0;
  }

  (async function(){
    var posted = false;
    try{
      /* ============ 0. 页面形态 ============ */
      sec('0. 页面形态（应完全没有内置数据）');
      log('window.WB_CFG（应为 null＝本机模式）', JSON.stringify(window.WB_CFG || null));
      log('HAS_BUILTIN', HAS_BUILTIN);
      log('GROUPS 分组数', GROUPS.length);
      log('ALL 条数（页面载入时）', ALL.length);
      log('内置分组标签数（应 0）', document.querySelectorAll('.gt[data-grp]').length);
      log('内置数据来源说明是否已隐藏', vis(id('srcrow')) === false);
      log('页头文案', txt(id('sub')).slice(0, 78));
      log('页头含「内置 31 条」（应 false）', txt(id('sub')).indexOf('内置') >= 0);
      log('统计卡：', Array.prototype.map.call(document.querySelectorAll('#stats .stat'),
            function(s){ return txt(s.querySelector('.k')) + '=' + txt(s.querySelector('.v')); }).join(' | '));
      log('第一张统计卡标题', txt(document.querySelector('#stats .stat .k')));
      log('空图提示可见（无数据时应有）', vis(id('cempty')));
      log('空图提示文字', txt(id('cempty')));
      log('快速选择按钮数', document.querySelectorAll('#quick button').length);
      log('全选按钮在', !!id('sel-all'));

      /* ============ 1. 导入 3 列 ============ */
      sec('1. 导入 3 列（标签「贵金属」）');
      var n = await importXlsx('贵金属');
      log('本次导入条数', n);
      log('CUSTOM 条数', CUSTOM.length);
      log('自动全选条数 / ALL 条数', state.active.size + ' / ' + ALL.length);
      log('图表 dataset 数', (chart && chart.data) ? chart.data.datasets.length : -1);
      log('空图提示此时应隐藏', vis(id('cempty')) === false);
      log('页头点数', txt(id('ptn')));
      log('chip 区标签分组', Array.prototype.map.call(document.querySelectorAll('.gt[data-tag]'),
            function(e){ return e.dataset.tag; }).join(' / '));
      log('「＋ 上传 Excel」在', !!id('chip-add'));

      var ser = CUSTOM.filter(function(s){ return s.src !== 'overlay'; });
      var sNames = ser.map(function(s){ return s.n; });
      var sUnits = ser.map(function(s){ return s.u || ''; });
      var sCnts  = ser.map(function(s){ return s.pts.length; });

      /* ============ 2. 备份面板 ============ */
      sec('2. 备份面板：Excel / CSV / json 三按钮');
      clickId('btn-bak');
      await waitFor(function(){ return !!id('bk-xlsx'); }, 4000, '打开备份面板');
      log('Excel 按钮存在', !!id('bk-xlsx'));
      log('CSV 按钮存在', !!id('bk-csv'));
      log('json 备份按钮存在', !!id('bk-out'));
      log('按钮文案', ['bk-xlsx','bk-csv','bk-out','bk-in','bk-clear'].map(function(k){
            return k + '=' + txt(id(k)); }).join(' ｜ '));
      log('说明文案提到宽表', txt(document.querySelector('#dlg .sub2')).indexOf('宽表') >= 0);

      /* 拦下下载：headless 里没必要真的落盘，顺便统计文件名与体积 */
      var dls = [];
      var realDownload = window.download;
      window.download = function(name, data, mime){
        dls.push({ name: name, size: (data && data.length) || 0, mime: mime || '' });
      };
      clickId('bk-xlsx'); await sleep(600);
      clickId('bk-csv');  await sleep(400);
      log('拦到的下载次数（应 2）', dls.length);
      log('Excel 下载', dls[0] ? (dls[0].name + ' · ' + dls[0].size + ' 字节 · ' + dls[0].mime) : '(无)');
      log('CSV 下载', dls[1] ? (dls[1].name + ' · ' + dls[1].size + ' 字节 · ' + dls[1].mime) : '(无)');
      log('Excel 文件名以 .xlsx 结尾', !!dls[0] && /\.xlsx$/.test(dls[0].name));
      log('CSV 文件名以 .csv 结尾', !!dls[1] && /\.csv$/.test(dls[1].name));
      window.download = realDownload;

      /* ============ 3. 导出的 xlsx 结构自检 ============ */
      sec('3. 导出的 xlsx 结构（页内自检）');
      var bytes = xlsxBytes(ser);
      log('xlsx 字节数', bytes ? bytes.length : 0);
      log('zip 魔数 PK', bytes && bytes[0] === 0x50 && bytes[1] === 0x4b);
      var w = wideTable(ser);
      log('宽表天数（日期并集）', w.days.length);
      log('宽表首末日期', dayToStr(w.days[0]) + ' → ' + dayToStr(w.days[w.days.length - 1]));
      log('宽表列数', w.cols.length + '（源 ' + ser.length + ' 条）');
      log('各列非空格子数', w.cols.map(function(c){ return c.filter(function(v){ return v !== null; }).length; }).join(' / '));
      log('与源点数一致', w.cols.map(function(c){ return c.filter(function(v){ return v !== null; }).length; }).join('|') === sCnts.join('|'));

      /* ============ 4. 往返：再拖回上传面板 ============ */
      sec('4. 导出的 xlsx 拖回上传面板（往返闭合）');
      clickId('btn-up');
      await waitFor(function(){ return !!id('fpick'); }, 4000, '打开上传面板');
      readUpFile(new File([bytes], '我的商品-往返.xlsx'));
      await waitFor(function(){ return UP.list && UP.list.length; }, 25000, '解析往返 xlsx');
      log('识别出的数值列数', UP.list.length + '（源 ' + ser.length + '）');
      log('列名', UP.list.map(function(it){ return it.name; }).join(' / '));
      log('单位', UP.list.map(function(it){ return it.unit || '(空)'; }).join(' / '));
      log('各列点数', UP.list.map(function(it){ return it.pts.length; }).join(' / '));
      log('列名与源逐列一致', UP.list.map(function(it){ return it.name; }).join('|') === sNames.join('|'));
      log('单位与源逐列一致', UP.list.map(function(it){ return it.unit || ''; }).join('|') === sUnits.join('|'));
      log('点数与源逐列一致', UP.list.map(function(it){ return it.pts.length; }).join('|') === sCnts.join('|'));
      log('首列首末日期', UP.list[0].pts.length
          ? dayToStr(UP.list[0].pts[0][0]) + ' → ' + dayToStr(UP.list[0].pts[UP.list[0].pts.length - 1][0]) : '—');
      log('首列首个价格', UP.list[0].pts.length ? UP.list[0].pts[0][1] : '—');
      log('源首列首个价格', ser[0] ? ser[0].pts[0][1] : '—');
      log('价格逐点一致', (function(){
            var a = UP.list[0].pts.slice(), b = ser[0].pts.slice();
            if(a.length !== b.length) return false;
            for(var i = 0; i < a.length; i++) if(a[i][0] !== b[i][0] || Math.abs(a[i][1] - b[i][1]) > 1e-9) return false;
            return true;
          })());
      closeDlg();
      await waitFor(function(){ return !id('mask').classList.contains('on'); }, 2000, '关上传面板');

      /* ============ 5. CSV 文本形状 ============ */
      sec('5. CSV 文本形状');
      var csv = csvWideText(ser);
      log('CSV 字符数', csv ? csv.length : 0);
      log('第一个字符是 BOM', !!csv && csv.charCodeAt(0) === 0xFEFF);
      var lines = csv.replace(/^\uFEFF/, '').split('\r\n').filter(function(l){ return l !== ''; });
      log('CSV 行数（应 = 2 表头 + 天数）', lines.length);
      log('第一行前 46 字', lines[0].slice(0, 46));
      log('第二行前 46 字', lines[1].slice(0, 46));
      log('第一行以「日期」开头', lines[0].indexOf('日期,') === 0);
      log('第二行以「单位」开头', lines[1].indexOf('单位,') === 0);
      log('第三行以日期开头', /^\d{4}-\d{2}-\d{2},/.test(lines[2]));
      log('第三行', lines[2].slice(0, 60));
      log('末行', lines[lines.length - 1].slice(0, 60));

      /* ============ 6. 把 xlsx 字节回传 Python（独立校验） ============ */
      sec('6. xlsx 字节回传 Python 用 openpyxl 独立核验');
      await fetch('/__xlsx__', { method: 'POST', body: JSON.stringify({
        b64: b64(bytes), names: sNames, units: sUnits, counts: sCnts,
        days: w.days.length, first: dayToStr(w.days[0]),
        last: dayToStr(w.days[w.days.length - 1])
      })});
      log('已回传字节数（base64 长度）', b64(bytes).length);

      /* ============ 7. 清空后空图提示回来 ============ */
      sec('7. 清空全部商品 → 空图提示回来');
      openBak();                                   // 第 4 段把弹窗换成了上传面板，这里重新打开
      await waitFor(function(){ return !!id('bk-clear'); }, 4000, '再开备份面板');
      clickId('bk-clear');
      await waitFor(function(){ return CUSTOM.length === 0; }, 4000, '清空');
      log('CUSTOM 已清零', CUSTOM.length);
      log('ALL 条数', ALL.length);
      log('空图提示可见', vis(id('cempty')));
      log('页头点数', txt(id('ptn')));
      log('chip 区只剩上传按钮', !document.querySelector('.chip.own') && !!id('chip-add'));

      out.push('');
      out.push('JS错误: ' + (errs.length ? errs.join(' | ') : '无'));
    }catch(e){
      out.push('✗ 探针中断：' + (e && e.message || e));
      out.push('JS错误: ' + (errs.length ? errs.join(' | ') : '无'));
    }
    await fetch('/__res__', { method: 'POST', body: out.join('\n') + '\n' });
    posted = true;
  })();
})();
</script>
'''


class Quiet(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def do_POST(self):
        n = int(self.headers.get('Content-Length') or 0)
        body = self.rfile.read(n).decode('utf-8', 'ignore')
        if self.path.startswith('/__xlsx__'):
            XLSXB64['body'] = body
        else:
            RESULT.append(body)
        self.send_response(200)
        self.send_header('Content-Length', '2')
        self.end_headers()
        self.wfile.write(b'ok')


def serve():
    h = functools.partial(Quiet, directory=SITE)
    socketserver.TCPServer.allow_reuse_address = True
    httpd = socketserver.TCPServer(('127.0.0.1', PORT), h)
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    return httpd


def run_probe(url, timeout=200):
    del RESULT[:]
    XLSXB64.clear()
    cmd = [EDGE, '--headless=new', '--no-sandbox', '--disable-gpu', '--disable-dev-shm-usage',
           '--window-size=1512,1600', url]
    p = subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    t0 = time.time()
    while time.time() - t0 < timeout:
        if RESULT and XLSXB64.get('body'):
            break
        time.sleep(0.25)
    time.sleep(0.4)
    try:
        p.terminate(); p.wait(timeout=8)
    except Exception:
        try: p.kill()
        except Exception: pass
    return (RESULT[0] if RESULT else '（%.0f 秒内没收到探针回传）' % timeout)


def check_xlsx(path):
    """Python 侧独立开箱：zip 完整性 + openpyxl 读结构 / 表头 / 日期 / 逐列非空格数"""
    import datetime
    import zipfile
    from openpyxl import load_workbook
    exp = json.loads(XLSXB64['body'])
    raw = base64.b64decode(exp['b64'])
    with open(path, 'wb') as f:
        f.write(raw)
    lines = []
    z = zipfile.ZipFile(path)
    broken = z.testzip()
    lines.append('zip 完整性（CRC 全过）      : %s' % ('✓ 通过' if broken is None else '✗ 损坏项 ' + str(broken)))
    lines.append('zip 条目                    : %s' % ' · '.join(z.namelist()))
    lines.append('文件体积                    : %.1f KB（%d 字节）' % (len(raw) / 1024, len(raw)))
    wb = load_workbook(path)
    lines.append('工作表                      : %s' % ' · '.join(wb.sheetnames))
    ok = True

    def chk(tag, got, want):
        nonlocal ok
        good = (got == want)
        ok = ok and good
        lines.append('%-26s: %s  (实测 %r / 期望 %r)' % (tag, '✓' if good else '✗', got, want))

    chk('工作表名', wb.sheetnames, ['日线数据', '品种说明'])
    ws = wb['日线数据']
    chk('A1 表头', ws.cell(1, 1).value, '日期')
    chk('第 1 行品种名', [ws.cell(1, 2 + i).value for i in range(len(exp['names']))], exp['names'])
    chk('A2 单位行首格', ws.cell(2, 1).value, '单位')
    chk('第 2 行单位', [(ws.cell(2, 2 + i).value or '') for i in range(len(exp['names']))], exp['units'])
    d0 = ws.cell(3, 1).value
    dN = ws.cell(2 + exp['days'], 1).value
    d0s = d0.strftime('%Y-%m-%d') if isinstance(d0, (datetime.date, datetime.datetime)) else str(d0)
    dNs = dN.strftime('%Y-%m-%d') if isinstance(dN, (datetime.date, datetime.datetime)) else str(dN)
    chk('日期是真日期型', isinstance(d0, (datetime.date, datetime.datetime)), True)
    chk('首行日期', d0s, exp['first'])
    chk('末行日期', dNs, exp['last'])
    chk('数据行数', ws.max_row - 2, exp['days'])
    counts = []
    for i in range(len(exp['names'])):
        c = 0
        for r in range(3, ws.max_row + 1):
            if ws.cell(r, 2 + i).value is not None:
                c += 1
        counts.append(c)
    chk('各列非空格子数', counts, exp['counts'])
    ws2 = wb['品种说明']
    chk('说明表头', [ws2.cell(1, c).value for c in range(1, 9)],
        ['#', '品种', '标签', '单位', '起始', '截止', '数据点', '来源'])
    chk('说明表行数', ws2.max_row - 1, len(exp['names']))
    chk('说明表首个品种', ws2.cell(2, 2).value, exp['names'][0])
    lines.append('说明表首行                  : %s' % ' | '.join(
        str(ws2.cell(2, c).value) for c in range(1, 9)))
    lines.append('日线表第 3 行               : %s' % ' | '.join(
        str(ws.cell(3, c).value) for c in range(1, min(6, len(exp['names']) + 2))))
    lines.append('')
    lines.append('独立校验结论                : %s' % ('✓ 全部一致' if ok else '✗ 有不一致'))
    return '\n'.join(lines), ok


def main():
    shutil.rmtree(SITE, ignore_errors=True)
    os.makedirs(SITE, exist_ok=True)
    size, rows = make_xlsx()
    print('=' * 78)
    print('「页面无内置数据 + 备份面板导出 Excel/CSV」E2E')
    print('=' * 78)
    print('源 xlsx：%s（%d 字节 · %d 行 · 3 个数值列 + 1 备注列）' % (XLSX_NAME, size, rows))

    src = io.open(SRC_PAGE, encoding='utf-8').read()
    # ★ 本机模式：把注入的云端配置掏空 —— 保证探针一条记录都不写进生产表
    stripped, cnt = re.subn(r'window\.WB_CFG\s*=\s*\{[^}]*\};', 'window.WB_CFG = null;', src)
    assert cnt == 1, '未能掏空 WB_CFG（匹配 %d 次）' % cnt
    # 进页面先把本机缓存清掉，保证每次跑都是干净起点
    i = src.find('<body')
    j = src.find('>', i)
    pre = ('<script>try{localStorage.removeItem("wb_c100_custom_v1");'
           'localStorage.removeItem("wb_c100_bover_v1");}catch(e){}</script>')
    page = stripped[:j + 1] + pre + stripped[j + 1:]
    io.open(os.path.join(SITE, PAGE), 'w', encoding='utf-8').write(page)
    io.open(os.path.join(SITE, PROBE_PAGE), 'w', encoding='utf-8').write(
        page.replace('</body>', DBG.replace('@XLSX@', XLSX_NAME) + '</body>', 1))

    srv = serve()
    res = run_probe('http://localhost:%d/%s' % (PORT, PROBE_PAGE), timeout=260)
    print(res)
    print('')
    print('=' * 78)
    print('Python 侧独立核验（openpyxl，不走页面自己的解析器）')
    print('=' * 78)
    if XLSXB64.get('body'):
        rep, ok = check_xlsx(OUT_XLSX)
        print(rep)
    else:
        ok = False
        print('✗ 没收到回传的 xlsx 字节')
    print('')
    print('=' * 78)
    srv.shutdown()
    shutil.rmtree(SITE, ignore_errors=True)
    print('（已清理探针临时目录 %s；导出的样例留在 %s）' % (os.path.relpath(SITE, HERE), OUT_XLSX))
    bad = ('✗' in res) or ('JS错误: 无' not in res) or (not ok)
    print('结论：' + ('✗ 有用例失败' if bad else '✓ 全部通过（页面 8 段 + 独立核验）'))
    return 1 if bad else 0


if __name__ == '__main__':
    sys.exit(main())
