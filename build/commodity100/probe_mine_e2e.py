#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""《大宗商品日线对比图》「我的商品」功能端到端探针。

验证：上传 xlsx → 列自动识别 → 预览 → 入库 → chip 出现 → 进入图表数据集
     → 编辑（改名/改值/改日期/补点/批量粘贴/删点）→ 保存 → 持久化
     → 管理面板 → 二次确认删除 → 清理干净

★ 关键纪律：无头 + --virtual-time-budget 会把 setTimeout 瞬间跑完，
  而 xlsx 解压是真实异步工作 —— 所以全部异步等待必须用 requestAnimationFrame 按帧轮询。

用法：python3 probe_mine_e2e.py
"""
import functools
import http.server
import io
import os
import shutil
import socketserver
import subprocess
import sys
import threading
import time

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
DELIVER = os.path.join(ROOT, 'deliver')
SRC_PAGE = os.path.join(DELIVER, '大宗商品日线对比图-2005-2026.html')
# ★ 探针在自己的临时目录里跑，绝不碰 deploy_site/（那里必须与线上保持一致）
SITE = os.path.join(HERE, '_probe_site')
PAGE = 'commodity-daily.html'
PROBE_PAGE = '_e2e_mine.html'
XLSX = os.path.join(SITE, '_e2e_mine.xlsx')
PORT = 8793
EDGE = '/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge'


# ---------------------------------------------------------------- 造测试用 xlsx
def make_xlsx():
    import datetime
    from openpyxl import Workbook
    wb = Workbook()

    # 表 1：带干扰列（品种 / 涨跌幅 / 备注），日期是真实日期格式
    ws = wb.active
    ws.title = '有色金属现货'
    ws.append(['日期', '品种', '现货价(元/吨)', '涨跌幅%', '备注'])
    base = 68000
    rows = 31
    for i in range(rows):
        d = datetime.date(2024, 1, 2) + datetime.timedelta(days=i)
        ws.append([d, '铜', base + i * 137.5, -2.75 + i * 0.25, '测试数据' if i == 0 else ''])
    for r in range(2, rows + 2):
        ws.cell(row=r, column=1).number_format = 'yyyy-mm-dd'

    # 表 2：只有两列（日期 / 收盘价），无表头，验证「表头行」自动判断
    ws2 = wb.create_sheet('单列价格')
    for i in range(12):
        d = datetime.date(2023, 6, 1) + datetime.timedelta(days=i)
        ws2.append([d, 4000 + i * 17])
    for r in range(1, 13):
        ws2.cell(row=r, column=1).number_format = 'yyyy-mm-dd'

    # 表 3：跳日数据 —— 「数据不全」最典型的形态（4 个月只给了 8 天）
    ws3 = wb.create_sheet('跳日现货')
    ws3.append(['交易日', '收盘价'])
    gaps = ['2024-06-03', '2024-06-18', '2024-07-02', '2024-07-22',
            '2024-08-05', '2024-08-19', '2024-09-02', '2024-09-20']
    vals = [73100, 73450, 72980, 73820, 74010, 73660, 74230, 74510]
    for ds, v in zip(gaps, vals):
        y, m, dd = (int(x) for x in ds.split('-'))
        ws3.append([datetime.date(y, m, dd), v])
    for r in range(2, 10):
        ws3.cell(row=r, column=1).number_format = 'yyyy-mm-dd'

    # 表 4：早期历史 —— 起点早于内置数据（1996-10-14），验证时间边界会自动撑大
    ws4 = wb.create_sheet('早期历史')
    ws4.append(['日期', '价格'])
    early = []
    early.append([datetime.date(1990, 1, 15), 2000])
    early.append([datetime.date(1990, 7, 1), 2100])
    for y in range(1991, 1997):
        early.append([datetime.date(y, 1, 15), 2000 + (y - 1990) * 55])
    for row in early:
        ws4.append(row)
    for r in range(2, len(early) + 2):
        ws4.cell(row=r, column=1).number_format = 'yyyy-mm-dd'

    wb.save(XLSX)
    return os.path.getsize(XLSX)


# ---------------------------------------------------------------- 探针脚本
DBG = r'''
<pre id="dbg">探针未完成</pre>
<script>
(function(){
  var out = [], errs = [];
  window.addEventListener('error', function(e){ errs.push('ERR: ' + (e.message||e)); });
  window.addEventListener('unhandledrejection', function(e){ errs.push('REJ: ' + ((e.reason&&e.reason.message)||e.reason)); });
  function log(k, v){ out.push(k + ' = ' + v); }
  function sec(t){ out.push(''); out.push('=== ' + t + ' ==='); }
  function dump(){
    var p = document.getElementById('dbg');
    var text = out.join('\n') + '\n\nJS错误: ' + (errs.length ? errs.join(' | ') : '无');
    p.textContent = text;
    p.style.cssText = 'position:relative;z-index:999;background:#000;color:#0f0;padding:10px;font-size:11px;white-space:pre-wrap';
    /* ★ 结果直接回传给本地服务端：不再依赖 --virtual-time-budget 的 dump 时机，
       虚拟时间与真实异步（xlsx 解压）赛跑是不稳定的 —— 已经踩过一次。 */
    try{ fetch('/__result', { method:'POST', body: text }); }catch(e){}
    document.title = 'READY';
  }
  function waitFor(fn, max, label){
    return new Promise(function(res, rej){
      var n = 0, fr = 0;
      (function tick(){
        fr++;
        var v = null;
        try{ v = fn(); }catch(e){ return rej(e); }
        if(v) return res(v);
        if(++n > max) return rej(new Error('等待超时: ' + label));
        requestAnimationFrame(tick);
      })();
    });
  }
  function txt(el){ return el ? String(el.textContent).replace(/\s+/g, ' ').trim() : '(缺)'; }

  (async function(){
    try{
      sec('0. 初始状态');
      log('内置序列数 BUILTIN', BUILTIN);
      log('初始序列总数 ALL', ALL.length);
      log('初始自定义数 CUSTOM', CUSTOM.length);
      log('页头数据点', document.getElementById('ptn').textContent);
      log('「我的商品」分组存在', !!document.querySelector('#chips .grp:last-child .gt'));
      log('分组名', txt(document.querySelector('#chips .grp:last-child .gt')));
      log('＋上传按钮存在', !!document.getElementById('chip-add'));
      log('toolbar 上传按钮', !!document.getElementById('btn-up'));

      sec('1. 上传 xlsx');
      var ab = await (await fetch('_e2e_mine.xlsx')).arrayBuffer();
      log('xlsx 字节', ab.byteLength);
      openUpload();
      log('弹窗已打开', document.getElementById('mask').classList.contains('on'));
      log('投放区存在', !!document.getElementById('drop'));
      readUpFile(new File([ab], '_e2e_mine.xlsx'));
      await waitFor(function(){ return UP.res && UP.res.pts.length >= 2; }, 4000, 'xlsx 解析');
      log('工作表', UP.names.length + ' 张 [' + UP.names.join(' / ') + ']');
      log('自动识别 日期列/价格列/表头行', UP.map.dateCol + ' / ' + UP.map.valCol + ' / ' + UP.map.hdr);
      log('解析点数', UP.res.pts.length);
      log('首点', dayToStr(UP.res.pts[0][0]) + ' → ' + UP.res.pts[0][1]);
      log('末点', dayToStr(UP.res.pts[UP.res.pts.length-1][0]) + ' → ' + UP.res.pts[UP.res.pts.length-1][1]);
      log('跳过行数', UP.res.bad);
      log('预览表行数', document.querySelectorAll('#up-prev tr').length);
      log('预览高亮列数', document.querySelectorAll('#up-prev .col-hi').length);
      log('解析提示', txt(document.getElementById('up-stat')));
      log('「进入对比」可点', !document.getElementById('up-ok').disabled);

      sec('2. 切工作表（两列无表头）');
      var sel = document.getElementById('up-sheet');
      sel.onchange({ target: { value: UP.names[1] } });
      await waitFor(function(){ return UP.sheet === UP.names[1] && UP.res && UP.res.pts.length >= 2; }, 3000, '切表');
      log('当前工作表', UP.sheet);
      log('识别 日期列/价格列/表头行', UP.map.dateCol + ' / ' + UP.map.valCol + ' / ' + UP.map.hdr);
      log('解析点数', UP.res.pts.length);
      log('首点', dayToStr(UP.res.pts[0][0]) + ' → ' + UP.res.pts[0][1]);

      sec('2.5 跳日数据（数据不全的典型形态）');
      document.getElementById('up-sheet').onchange({ target: { value: '跳日现货' } });
      await waitFor(function(){ return UP.sheet === '跳日现货' && UP.res && UP.res.pts.length === 8; }, 3000, '跳日表');
      log('识别 日期列/价格列', UP.map.dateCol + ' / ' + UP.map.valCol);
      log('解析点数（表里只有 8 天）', UP.res.pts.length);
      log('首点', dayToStr(UP.res.pts[0][0]) + ' → ' + UP.res.pts[0][1]);
      log('末点', dayToStr(UP.res.pts[7][0]) + ' → ' + UP.res.pts[7][1]);
      log('跨度 vs 点数', (UP.res.pts[7][0] - UP.res.pts[0][0]) + ' 天 / ' + UP.res.pts.length + ' 点（缺口可后续在编辑里补）');

      sec('2.6 早期历史（起点早于内置数据）');
      document.getElementById('up-sheet').onchange({ target: { value: '早期历史' } });
      await waitFor(function(){ return UP.sheet === '早期历史' && UP.res && UP.res.pts.length >= 8; }, 3000, '早期表');
      log('解析点数', UP.res.pts.length);
      log('最早点', dayToStr(UP.res.pts[0][0]));
      log('入库前 FIRST', dayToStr(FIRST));

      sec('3. 切回表1 并入库');
      document.getElementById('up-sheet').onchange({ target: { value: UP.names[0] } });
      await waitFor(function(){ return UP.sheet === UP.names[0] && UP.res && UP.res.pts.length >= 2; }, 3000, '切回');
      document.getElementById('up-name').value = 'E2E 测试商品';
      document.getElementById('up-unit').value = '元/吨';
      commitUpload();
      await waitFor(function(){ return CUSTOM.length === 1; }, 400, '入库');
      var s = CUSTOM[0];
      log('入库成功 自定义数', CUSTOM.length);
      log('序列总数 ALL', ALL.length);
      log('BYCODE 命中', !!BYCODE[s.c]);
      log('自动勾选', state.active.has(s.c));
      log('chip 存在', !!document.querySelector('.chip[data-c="' + s.c + '"]'));
      log('chip 文本', txt(document.querySelector('.chip[data-c="' + s.c + '"]')));
      log('chip 有编辑角标', !!document.querySelector('.chip[data-c="' + s.c + '"] .ed'));
      log('弹窗已关', !document.getElementById('mask').classList.contains('on'));
      log('页头点数', document.getElementById('ptn').textContent);
      log('统计条', txt(document.getElementById('stats')).slice(0, 150));

      /* 再传一条，确认自动配色不会和已有曲线撞色 */
      var ab2c = await (await fetch('_e2e_mine.xlsx')).arrayBuffer();
      openUpload();
      readUpFile(new File([ab2c], '_e2e_mine.xlsx'));
      await waitFor(function(){ return UP.res && UP.res.pts.length >= 2; }, 4000, '再次解析');
      var to2 = UP.names[1];
      document.getElementById('up-sheet').onchange({ target: { value: to2 } });
      await waitFor(function(){ return UP.sheet === to2 && UP.res && UP.res.pts.length >= 2; }, 3000, '再切表');
      document.getElementById('up-name').value = 'E2E 第二条';
      commitUpload();
      await waitFor(function(){ return CUSTOM.length === 2; }, 400, '第二条');
      log('两条自定义配色', CUSTOM[0].col + ' / ' + CUSTOM[1].col);
      log('配色不重复', CUSTOM[0].col !== CUSTOM[1].col);
      log('两条都在图里', chart.data.datasets.filter(function(d){ return d._own; }).length);
      var c2 = CUSTOM[1].c;
      CUSTOM = CUSTOM.filter(function(x){ return x.c !== c2; });
      state.active.delete(c2);
      afterCustomChange();
      log('清理回 1 条', CUSTOM.length);

      sec('4. 进入图表数据集');
      redraw();
      var labels = chart.data.datasets.map(function(d){ return d.label; });
      log('图表序列数', labels.length);
      log('含自定义曲线', labels.indexOf('E2E 测试商品') >= 0);
      log('带 _own 标记的数据集数', chart.data.datasets.filter(function(d){ return d._own; }).length);
      var ds = chart.data.datasets.filter(function(d){ return d._c === s.c; })[0];
      log('自定义曲线点数', ds ? ds.data.length : -1);
      log('曲线首点', ds ? dayToStr(ds.data[0].x) + ' = ' + ds.data[0].y : '-');
      log('曲线末点', ds ? dayToStr(ds.data[ds.data.length-1].x) + ' = ' + ds.data[ds.data.length-1].y : '-');
      log('归一化后首点应为100', ds ? ds.data[0].y : '-');

      sec('5. 分组快速选择 / 全选不误伤');
      var before = state.active.size;
      document.getElementById('sel-all').onclick();
      log('全选后已选数 (应=ALL)', state.active.size + ' / ' + ALL.length);
      document.getElementById('sel-all').onclick();
      log('再点全不选', state.active.size);
      state.active.add(s.c); buildChips(); redraw();
      log('重新勾选自定义', state.active.has(s.c) && chart.data.datasets.some(function(d){ return d._c === s.c; }));

      sec('6. 编辑：改名 / 改值 / 改日期');
      openEdit(s.c);
      log('面板已开', document.getElementById('mask').classList.contains('on'));
      log('标题', txt(document.querySelector('#dlg h3')));
      log('数据点输入框数', document.querySelectorAll('#dlg input[data-i]').length);
      log('分页信息', txt(document.querySelector('#dlg .pager')));
      var oldV = ED.pts[0][1], oldD = ED.pts[0][0];
      var inv = document.querySelector('#dlg input[data-f="v"]');
      inv.value = '99999'; inv.onchange();
      log('改价格 ' + oldV + ' → ' + ED.pts[0][1], ED.pts[0][1] === 99999);
      var ind = document.querySelector('#dlg input[data-f="d"]');
      ind.value = '2024-02-05'; ind.onchange();
      log('改日期 ' + dayToStr(oldD) + ' → ' + dayToStr(ED.pts[0][0]), dayToStr(ED.pts[0][0]) === '2024-02-05');
      log('脏标记', txt(document.getElementById('ed-note')));

      sec('7. 补点 / 批量粘贴 / 删点');
      document.getElementById('ed-nd').value = '2024-03-01';
      document.getElementById('ed-nv').value = '52000';
      document.getElementById('ed-add').onclick();
      log('补点后点数', ED.pts.length);
      log('末点', dayToStr(ED.pts[ED.pts.length-1][0]) + ' → ' + ED.pts[ED.pts.length-1][1]);
      // 覆盖已有日期
      document.getElementById('ed-nd').value = '2024-03-01';
      document.getElementById('ed-nv').value = '55555';
      document.getElementById('ed-add').onclick();
      log('同日再写应覆盖 (点数不变)', ED.pts.length + ' · 值 ' + ED.pts[ED.pts.length-1][1]);
      // 批量粘贴（含 1 行垃圾）
      document.getElementById('ed-paste').value =
        '2024-03-04,53100\n2024-03-05,53300\n这不是数据\n2024-03-06,53500';
      document.getElementById('ed-merge').onclick();
      log('粘贴后点数', ED.pts.length);
      log('粘贴提示', txt(document.getElementById('ed-note')));
      var n0 = ED.pts.length;
      document.querySelector('#dlg button[data-del]').onclick();
      log('删点 ' + n0 + ' → ' + ED.pts.length, ED.pts.length === n0 - 1);

      sec('8. 保存并回写');
      document.getElementById('ed-name').value = 'E2E 改名商品';
      document.getElementById('ed-unit').value = '元/千克';
      document.getElementById('ed-save').onclick();
      await waitFor(function(){ return CUSTOM[0] && CUSTOM[0].n === 'E2E 改名商品'; }, 400, '保存');
      log('名称', CUSTOM[0].n);
      log('单位', CUSTOM[0].u);
      log('点数', CUSTOM[0].pts.length);
      log('首点已改', dayToStr(CUSTOM[0].pts[0][0]) + ' = ' + CUSTOM[0].pts[0][1]);
      log('弹窗已关', !document.getElementById('mask').classList.contains('on'));
      log('chip 文本已更新', txt(document.querySelector('.chip[data-c="' + s.c + '"]')));
      log('排行表含新名', document.getElementById('rank').textContent.indexOf('E2E 改名商品') >= 0);
      log('覆盖表含新名', document.getElementById('cover').textContent.indexOf('E2E 改名商品') >= 0);
      log('统计条含「自定义 1 条」', document.getElementById('stats').textContent.indexOf('自定义 1 条') >= 0);
      var lbl2 = chart.data.datasets.map(function(d){ return d.label; });
      log('图表标签已更新', lbl2.indexOf('E2E 改名商品') >= 0);

      sec('9. 持久化与导出');
      var raw = localStorage.getItem(CUSTOM_KEY) || '';
      log('localStorage 长度', raw.length);
      log('localStorage 含新名', raw.indexOf('E2E 改名商品') >= 0);
      var csv = csvOf(BYCODE[s.c]);
      log('CSV 表头', csv.split('\n')[0]);
      log('CSV 行数(含表头)', csv.trim().split('\n').length);
      log('CSV 第2行', csv.split('\n')[1]);

      sec('10. 管理面板');
      openMgr();
      log('面板已开', document.getElementById('mask').classList.contains('on'));
      log('列表行数', document.querySelectorAll('#dlg table.ftb tr').length - 1);
      log('含商品名', document.getElementById('dlg').textContent.indexOf('E2E 改名商品') >= 0);
      log('含点数', document.getElementById('dlg').textContent.indexOf(String(CUSTOM[0].pts.length)) >= 0);
      log('有 隐藏/编辑/删除 按钮',
          !!document.querySelector('#dlg button[data-tg]') &&
          !!document.querySelector('#dlg button[data-ed]') &&
          !!document.querySelector('#dlg button[data-rm]'));
      document.querySelector('#dlg button[data-tg]').onclick();
      log('点隐藏后 active', state.active.has(s.c));
      document.querySelector('#dlg button[data-tg]').onclick();
      log('再点显示后 active', state.active.has(s.c));

      sec('11. 备份面板');
      openBak();
      log('面板已开', document.getElementById('mask').classList.contains('on'));
      log('列出条目', document.getElementById('dlg').textContent.indexOf('E2E 改名商品') >= 0);
      log('有导出/导入/清空', !!document.getElementById('bk-out') && !!document.getElementById('bk-in') && !!document.getElementById('bk-clear'));
      var bak = JSON.stringify({ v: 1, series: CUSTOM });
      log('备份 JSON 可序列化', bak.length + ' 字符');
      log('备份含点数', JSON.parse(bak).series[0].pts.length);
      closeDlg();

      sec('12. 二次确认删除');
      openEdit(s.c);
      var delBtn = document.getElementById('ed-del');
      delBtn.onclick();
      log('首次点删除 文案', txt(delBtn));
      log('仍未删除', CUSTOM.length === 1);
      delBtn.onclick();
      await waitFor(function(){ return CUSTOM.length === 0; }, 400, '删除');
      log('删除后自定义数', CUSTOM.length);
      log('BYCODE 已清', !BYCODE[s.c]);
      log('ALL 已移除', !ALL.some(function(x){ return x.c === s.c; }));
      log('active 已清', !state.active.has(s.c));
      log('图表已移除', !chart.data.datasets.some(function(d){ return d._c === s.c; }));
      log('chip 已消失', !document.querySelector('.chip[data-c="' + s.c + '"]'));
      log('localStorage 清空', localStorage.getItem(CUSTOM_KEY) === '[]');
      log('统计条回到「内置」', document.getElementById('stats').textContent.indexOf('自定义') < 0);
      log('页头点数', document.getElementById('ptn').textContent);

      sec('13. 时间边界自动扩展 / 缩回（上传 1990 年的数据）');
      log('入库前 FIRST', dayToStr(FIRST));
      var ab2 = await (await fetch('_e2e_mine.xlsx')).arrayBuffer();
      openUpload();
      readUpFile(new File([ab2], '_e2e_mine.xlsx'));
      await waitFor(function(){ return UP.res && UP.res.pts.length >= 2; }, 4000, '二次解析');
      document.getElementById('up-sheet').onchange({ target: { value: '早期历史' } });
      await waitFor(function(){ return UP.sheet === '早期历史' && UP.res && UP.res.pts.length >= 8; }, 3000, '早期表');
      log('解析点数', UP.res.pts.length);
      log('最早点', dayToStr(UP.res.pts[0][0]));
      document.getElementById('up-name').value = 'E2E 早期历史';
      commitUpload();
      await waitFor(function(){ return CUSTOM.length === 1; }, 400, '入库早期');
      log('入库后 FIRST（应提前到 1990）', dayToStr(FIRST));
      log('时间跨度(年)', ((LAST - FIRST) / 365.25).toFixed(1));
      log('日期控件 min 已更新', document.getElementById('dt-from').min);
      log('当前视窗（未被强行打断）', dayToStr(Math.round(state.view.min)) + ' → ' + dayToStr(Math.round(state.view.max)));
      log('窄视窗下曲线不在图内', !chart.data.datasets.some(function(d){ return d._c === CUSTOM[0].c; }));
      setRangeDays(null, 'all');
      log('切「全部」后曲线进入图表', chart.data.datasets.some(function(d){ return d._c === CUSTOM[0].c; }));
      log('图表标签', chart.data.datasets.filter(function(d){ return d._own; }).map(function(d){ return d.label; }).join(','));
      CUSTOM = [];
      afterCustomChange();
      log('清空后 FIRST 缩回', dayToStr(FIRST));
      log('清空后时间跨度(年)', ((LAST - FIRST) / 365.25).toFixed(1));
      log('BYCODE 无孤儿键', !Object.keys(BYCODE).some(function(c){ return c.indexOf('u_') === 0; }));
      log('ALL 只剩内置', ALL.length + ' / ' + BUILTIN);
      log('localStorage 干净', localStorage.getItem(CUSTOM_KEY) === '[]');
    }catch(e){
      log('!!! 流程异常', (e && e.message) || e);
      if(e && e.stack) out.push(String(e.stack).split('\n').slice(0, 4).join(' | '));
    }
    dump();
  })();
})();
</script>
'''


RESULT = []


class Quiet(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def do_POST(self):
        n = int(self.headers.get('Content-Length') or 0)
        RESULT.append(self.rfile.read(n).decode('utf-8', 'ignore'))
        self.send_response(200)
        self.send_header('Content-Length', '2')
        self.end_headers()
        self.wfile.write(b'ok')


def serve():
    h = functools.partial(Quiet, directory=SITE)
    socketserver.TCPServer.allow_reuse_address = True
    httpd = socketserver.TCPServer(('127.0.0.1', PORT), h)
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    return httpd


def run_probe(url, timeout=240):
    """跑一次无头浏览器，等页内探针把结果 POST 回来（真实等待，不用虚拟时间）。"""
    del RESULT[:]
    cmd = [EDGE, '--headless=new', '--no-sandbox', '--disable-gpu', '--disable-dev-shm-usage',
           '--window-size=1512,1600', url]
    p = subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    t0 = time.time()
    while time.time() - t0 < timeout:
        if RESULT:
            break
        time.sleep(0.25)
    time.sleep(0.4)
    try:
        p.terminate(); p.wait(timeout=8)
    except Exception:
        try: p.kill()
        except Exception: pass
    if RESULT:
        return RESULT[0]
    return '（%.0f 秒内没收到探针回传 —— 页面可能在早期就抛错了）' % timeout


def main():
    shutil.rmtree(SITE, ignore_errors=True)
    os.makedirs(SITE, exist_ok=True)
    n = make_xlsx()
    print('测试用 xlsx：%s（%d 字节）' % (os.path.basename(XLSX), n))
    src = io.open(SRC_PAGE, encoding='utf-8').read()
    io.open(os.path.join(SITE, PAGE), 'w', encoding='utf-8').write(src)
    io.open(os.path.join(SITE, PROBE_PAGE), 'w', encoding='utf-8').write(
        src.replace('</body>', DBG + '</body>', 1))
    srv = serve()
    print('\n' + '=' * 72)
    print(run_probe('http://localhost:%d/%s' % (PORT, PROBE_PAGE)))
    print('=' * 72)
    srv.shutdown()
    shutil.rmtree(SITE, ignore_errors=True)
    print('\n（已清理探针临时目录 %s）' % os.path.relpath(SITE, HERE))
    return 0


if __name__ == '__main__':
    sys.exit(main())
