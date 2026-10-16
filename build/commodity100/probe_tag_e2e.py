#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""「商品标签（自定义分组）+ 管理面板批量删除/全选」E2E 探针。

★ 本探针只测 UI 与分组逻辑，**必须跑在本机模式**：
   main() 里把注入的 window.WB_CFG 掏空 → WBCloud.off() 为 true → 所有写入退回本机，
   一条生产数据都不会落库（wb_c100_series 保持干净）。

覆盖：
  0. 初始结构：4 个内置分组 + 「＋ 上传 Excel」
  1. 上传面板有「标签」字段（可选已有标签，也能敲新名字）→ 真实点击「全部导入」
  2. chip 区按标签各自成组（不再全挤在「我的商品」一组）
  3. 编辑面板改标签 → 商品当场搬到另一组；与内置分类同名时补「我的」小标
  4. 点分组名＝整组选中 / 取消
  5. 标签留空 → 归入「我的商品」组，且排在最后
  6. 管理面板：全选复选框 + 已选计数 + 批量删除（真实点击，confirm 自动确认）
  7. 管理面板：按标签快捷勾选 + 批量改标签（prompt 打桩）
  8. 内置覆盖记录：标签列显示「—」，批量删除＝恢复内置原始数据
  9. 全程零 JS 错误

跑法：
  /Users/luoshui/.workbuddy/binaries/python/envs/default/bin/python probe_tag_e2e.py
"""
import functools
import http.server
import io
import os
import re
import shutil
import socketserver
import subprocess
import sys
import threading
import time

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
SRC_PAGE = os.path.join(ROOT, 'deliver', '大宗商品日线对比图-2005-2026.html')
SITE = os.path.join(HERE, '_tag_site')            # 独立临时目录：绝不碰 deploy_site/
PAGE = 'commodity-daily.html'
PROBE_PAGE = '_e2e_tag.html'
XLSX_NAME = '_tag_wide.xlsx'
XLSX = os.path.join(SITE, XLSX_NAME)
PORT = 8804
EDGE = '/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge'
RESULT = []


def make_xlsx():
    """三个数值列 + 一列备注，最后一列是纯文本（不该被当数值列）"""
    import datetime
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = '日线'
    ws.append(['日期', '沪金', '沪银', '沪铜', '备注'])
    ws.append(['单位', '元/克', '元/千克', '元/吨', ''])
    d = datetime.date(2024, 1, 1)
    i = 0
    while i < 90:
        if d.weekday() < 5:
            ws.append([d, round(480 + i * 0.9, 2), round(5900 + i * 11.3, 1),
                       round(74000 + i * 120.0, 1), '现货' if i % 9 == 0 else ''])
        d += datetime.timedelta(days=1)
        i += 1
    for r in range(3, ws.max_row + 1):
        ws.cell(row=r, column=1).number_format = 'yyyy-mm-dd'
    wb.save(XLSX)
    return os.path.getsize(XLSX), ws.max_row


DBG = r'''
<pre id="dbg" style="display:none"></pre>
<script>
(function(){
  var errs = [], out = [];
  window.addEventListener('error', function(e){ errs.push('ERR: ' + e.message); });
  window.addEventListener('unhandledrejection', function(e){
    errs.push('REJ: ' + ((e.reason && e.reason.message) || e.reason));
  });
  window.confirm = function(){ return true; };        // headless 下 confirm 会阻塞
  window.__PROMPT = '';
  window.prompt = function(){ return window.__PROMPT; };

  function log(k, v){ out.push(k + ' = ' + v); }
  /* 页面可能没有内置数据（build_opts.json → builtin:false）；内置相关用例据此跳过 */
  var HASB = (typeof HAS_BUILTIN === 'undefined') ? true : HAS_BUILTIN;
  function sec(t){ out.push(''); out.push('=== ' + t + ' ==='); }
  function id(i){ return document.getElementById(i); }
  function txt(el){ return el ? (el.textContent || '').replace(/\s+/g, ' ').trim() : '(null)'; }
  function click(el){ if(!el){ out.push('✗ 找不到元素'); return false; } el.click(); return true; }
  function clickId(i){ return click(id(i)); }
  function sleep(ms){ return new Promise(function(r){ setTimeout(r, ms); }); }
  function waitFor(fn, ms, tag){
    return new Promise(function(res, rej){
      var t0 = Date.now();
      (function poll(){
        var v; try{ v = fn(); }catch(e){ return rej(new Error(tag + ' 抛错: ' + e.message)); }
        if(v) return res(v);
        if(Date.now() - t0 > (ms || 3000)) return rej(new Error(tag + ' 超时(' + (ms || 3000) + 'ms)'));
        requestAnimationFrame(poll);
      })();
    });
  }
  function tags(){ return Array.from(document.querySelectorAll('.gt[data-tag]')).map(function(e){ return e.dataset.tag; }); }
  function grpCount(t){ var g = document.querySelector('.gt[data-tag="' + t + '"]'); return g ? g.parentNode.querySelectorAll('.chip[data-c]').length : -1; }
  function setVal(el, v){
    el.value = v;
    el.dispatchEvent(new Event('input',  { bubbles: true }));
    el.dispatchEvent(new Event('change', { bubbles: true }));
  }
  async function importXlsx(tag){
    clickId('btn-up');
    await waitFor(function(){ return !!id('fpick'); }, 4000, '打开上传面板');
    var ab = await (await fetch('@XLSX@')).arrayBuffer();
    readUpFile(new File([ab], '标签测试.xlsx'));
    await waitFor(function(){ return UP.list && UP.list.length >= 3; }, 20000, '解析 xlsx');
    var n0 = CUSTOM.length;
    if(tag !== null && id('up-tag')) setVal(id('up-tag'), tag);
    clickId('up-ok');
    await waitFor(function(){ return CUSTOM.length > n0; }, 6000, '导入落地');
    return CUSTOM.length - n0;
  }

  (async function(){
    try{
      /* ============ 0. 初始结构 ============ */
      sec('0. 初始结构');
      log('云端配置（应为空＝本机模式）', JSON.stringify(window.WB_CFG || null));
      log('CUSTOM 条数', CUSTOM.length);
      log('内置分组标签', Array.from(document.querySelectorAll('.gt[data-grp]')).map(function(e){ return e.dataset.grp; }).join(' / '));
      log('自建标签分组', tags().join(' / ') || '（无）');
      log('「＋ 上传 Excel」按钮在', !!id('chip-add'));

      /* ============ 1. 上传面板的「标签」字段 ============ */
      sec('1. 上传面板：标签字段');
      clickId('btn-up');
      await waitFor(function(){ return !!id('fpick'); }, 4000, '打开上传面板');
      /* ★ 标签字段在「解析出表格之后」才渲染（面板先显示拖拽区），所以要先喂一个文件 */
      var ab0 = await (await fetch('@XLSX@')).arrayBuffer();
      readUpFile(new File([ab0], '标签测试.xlsx'));
      await waitFor(function(){ return UP.list && UP.list.length >= 3; }, 20000, '解析 xlsx');
      log('标签输入框存在', !!id('up-tag'));
      log('标签字段默认值', id('up-tag') ? id('up-tag').value : '(无)');
      var dl = id('up-tag-dl');
      log('可选标签（datalist）', dl ? Array.from(dl.options).map(function(o){ return o.value; }).join(' / ') : '(无)');
      log('说明文案含「单独选」（分组列说明）', document.getElementById('upbody').textContent.indexOf('单独选') >= 0);
      log('默认分组输入框旁有「套用到所有行」', !!id('up-tagall'));
      log('每行都有分组下拉', document.querySelectorAll('#up-list select[data-tg]').length === UP.list.length);
      var keepOk = (function(){
        setVal(id('up-tag'), '钼系');
        var sh = id('up-sheet');
        sh.onchange({ target: { value: UP.names[0] } });   // 换工作表会整表重渲染，标签不该丢
        return !!id('up-tag') && id('up-tag').value === '钼系';
      })();
      log('换工作表后标签不丢', keepOk);
      log('数值列识别数（备注列不该被算进来）', UP.list.length);
      closeDlg();
      await waitFor(function(){ return !id('mask').classList.contains('on'); }, 2000, '关面板');

      /* ============ 2. 带标签导入 → 按标签成组 ============ */
      sec('2. 导入 3 列并打标签「钼系」');
      var n = await importXlsx('钼系');
      log('本次导入条数', n);
      log('3 条都带标签 钼系', CUSTOM.filter(function(s){ return s.tag === '钼系'; }).length);
      log('chip 区标签分组', tags().join(' / '));
      log('「钼系」组里 chip 数', grpCount('钼系'));
      log('虚线 chip（我的商品）数', document.querySelectorAll('.chip.own').length);
      log('第一组自建标签带分隔线', HASB
          ? document.querySelector('.gt[data-tag]').parentNode.classList.contains('sep')
          : '⊘ 本页无内置数据，自建标签就是第一组，不画分隔线');
      log('chip 气泡含标签名', document.querySelector('.chip.own').title.indexOf('钼系') >= 0);

      /* ============ 3. 编辑面板改标签 ============ */
      sec('3. 编辑面板改标签 → 商品当场换组');
      var first = CUSTOM[0];
      click(document.querySelector('.chip[data-c="' + first.c + '"] .ed'));
      await waitFor(function(){ return !!id('ed-tag'); }, 4000, '打开编辑面板');
      log('编辑面板有标签字段', !!id('ed-tag'));
      log('回显当前标签', id('ed-tag').value);
      setVal(id('ed-tag'), '有色金属');           // 与内置分类同名
      clickId('ed-save');
      await waitFor(function(){ return !CUSTOM.some(function(s){ return s.c === first.c && s.tag !== '有色金属'; }); }, 4000, '标签保存');
      log('保存后该条标签', CUSTOM.filter(function(s){ return s.c === first.c; })[0].tag);
      log('分组顺序（同名标签走内置次序）', tags().join(' / '));
      log('「有色金属」组 chip 数', grpCount('有色金属'));
      log('「钼系」组 chip 数（应少 1）', grpCount('钼系'));
      log('同名标签带「我的」小标', HASB
          ? !!document.querySelector('.gt[data-tag="有色金属"] .mine')
          : '⊘ 本页无内置数据，不显示「我的」小标');
      log('内置分组标签仍 4 个', document.querySelectorAll('.gt[data-grp]').length);

      /* ============ 4. 点分组名＝整组选中 / 取消 ============ */
      sec('4. 点分组名＝整组选中 / 取消');
      var codes = CUSTOM.filter(function(s){ return s.tag === '钼系'; }).map(function(s){ return s.c; });
      codes.forEach(function(c){ state.active.delete(c); });   // 先置成「全不选」这个已知起点
      buildChips(); redraw();
      click(document.querySelector('.gt[data-tag="钼系"]'));
      log('起点全不选 → 点一下该组全选', codes.every(function(c){ return state.active.has(c); }));
      click(document.querySelector('.gt[data-tag="钼系"]'));
      log('再点一下 → 该组全不选', codes.every(function(c){ return !state.active.has(c); }));
      if(!HASB){
        log('内置组整组翻转', '⊘ 本页无内置数据，跳过');
      }else{
      var gc = GROUPS.filter(function(x){ return x.name === '农产品'; })[0].series.map(function(s){ return s.c; });
      var beforeAny = gc.filter(function(c){ return state.active.has(c); }).length;
      click(document.querySelector('.gt[data-grp="农产品"]'));
      var afterAny = gc.filter(function(c){ return state.active.has(c); }).length;
      log('内置组「农产品」选中数', beforeAny + ' → ' + afterAny + '（共 ' + gc.length + ' 条，应整组翻转）');
      log('内置组整组翻转正确', afterAny === (beforeAny === gc.length ? 0 : gc.length));
      click(document.querySelector('.gt[data-grp="农产品"]'));
      }

      /* ============ 5. 标签留空 → 我的商品组，且排最后 ============ */
      sec('5. 标签留空＝归入「我的商品」，排在最后');
      var n2 = await importXlsx('');
      log('本次导入条数', n2);
      log('归入「我的商品」组条数', CUSTOM.filter(function(s){ return tagOf(s) === TAG_DEFAULT; }).length);
      log('留空时库里存的是空串', CUSTOM.filter(function(s){ return !s.tag; }).length);
      log('分组顺序', tags().join(' / '));
      log('末组就是「我的商品」', tags()[tags().length - 1]);
      log('「＋ 上传 Excel」仍在', !!id('chip-add'));

      /* ============ 6. 管理面板：全选 + 批量删除 ============ */
      sec('6. 管理面板：全选 + 批量删除');
      clickId('btn-my');
      await waitFor(function(){ return !!id('mg-all'); }, 4000, '打开管理面板');
      var cks = function(){ return Array.from(document.querySelectorAll('#dlg input.mg-ck')); };
      log('列表行数 / 复选框数', (document.querySelectorAll('#dlg table.ftb tr').length - 1) + ' / ' + cks().length);
      log('初始未选中', cks().filter(function(c){ return c.checked; }).length);
      log('删除按钮初始 disabled', id('mg-del').disabled);
      click(id('mg-all'));
      log('点全选后已勾选', cks().filter(function(c){ return c.checked; }).length);
      log('已选计数文案', txt(id('mg-selinfo')));
      log('删除按钮文案', txt(id('mg-del')) + ' | disabled=' + id('mg-del').disabled);
      click(id('mg-all'));                                   // 取消全选
      log('取消全选后已勾选', cks().filter(function(c){ return c.checked; }).length);
      log('取消后删除按钮 disabled', id('mg-del').disabled);

      var delTag = '钼系';
      var delIds = CUSTOM.filter(function(s){ return s.tag === delTag; }).map(function(s){ return s.c; });
      click(document.querySelector('#dlg button[data-picktag="' + delTag + '"]'));
      log('按标签「' + delTag + '」勾选数', cks().filter(function(c){ return c.checked; }).length);
      var before = CUSTOM.length;
      clickId('mg-del');
      await waitFor(function(){ return CUSTOM.length < before; }, 4000, '批量删除');
      log('删除后 CUSTOM 条数', CUSTOM.length + '（原 ' + before + '）');
      log('该标签记录已清空', CUSTOM.filter(function(s){ return s.tag === delTag; }).length === 0);
      log('BYCODE 已清', delIds.every(function(c){ return !BYCODE[c]; }));
      log('chip 已消失', delIds.every(function(c){ return !document.querySelector('.chip[data-c="' + c + '"]'); }));
      log('该标签分组已消失', tags().indexOf(delTag) < 0);
      log('localStorage 已同步', localStorage.getItem('wb_c100_custom_v1').indexOf(delTag) < 0);

      /* ============ 7. 批量改标签 ============ */
      sec('7. 管理面板：批量改标签');
      click(id('mg-all'));
      window.__PROMPT = '钢铁原料';
      clickId('mg-tagset');
      await sleep(120);
      await waitFor(function(){ return CUSTOM.every(function(s){ return s.tag === '钢铁原料'; }); }, 4000, '批量改标签');
      log('全部记录的新标签', Array.from(new Set(CUSTOM.map(function(s){ return s.tag; }))).join(' / '));
      log('分组顺序', tags().join(' / '));
      log('该组 chip 数', grpCount('钢铁原料') + ' / CUSTOM ' + CUSTOM.length);
      log('管理面板标签列显示新标签', txt(document.querySelector('#dlg .mgtag')));

      /* ============ 8. 内置覆盖记录：标签列「—」，批量删除＝恢复原始 ============ */
      if(!HASB){
        sec('8. 内置覆盖记录');
        closeDlg();
        log('内置覆盖记录', '⊘ 本页无内置数据，跳过');
      }else{
      sec('8. 内置覆盖记录（ov_gold）');
      closeDlg();
      click(document.querySelector('.chip[data-c="gold"] .ed'));
      await waitFor(function(){ return !!id('eb-data'); }, 4000, '打开内置编辑面板');
      clickId('eb-data');
      await waitFor(function(){ return !!id('ed-save'); }, 4000, '进入数据编辑');
      clickId('ed-cancel');
      var gold0 = BYCODE['gold'].pts0.length;
      log('已生成覆盖记录', !!CUSTOM.filter(function(s){ return s.src === 'overlay'; }).length);
      clickId('btn-my');
      await waitFor(function(){ return !!id('mg-all'); }, 4000, '再开管理面板');
      var ovRow = document.querySelector('#dlg tr:has(.ovtag)');
      log('管理面板出现「内置覆盖」行', !!ovRow);
      log('该行标签列', ovRow ? txt(ovRow.children[2]) : '(无)');
      click(id('mg-all'));
      clickId('mg-del');
      await waitFor(function(){ return !CUSTOM.some(function(s){ return s.src === 'overlay'; }); }, 4000, '批量删除覆盖记录');
      log('覆盖记录已删', !CUSTOM.some(function(s){ return s.src === 'overlay'; }));
      log('内置 gold 已恢复原始数据', BYCODE['gold'].pts.length === gold0 && BYCODE['gold'].pts === BYCODE['gold'].pts0);
      log('再开管理面板无覆盖行', (function(){ openMgr(); return !document.querySelector('#dlg .ovtag'); })());
      }

      /* ============ 9. 收尾：清干净 + 报错 ============ */
      sec('9. 收尾');
      closeDlg();
      CUSTOM.slice().forEach(function(s){
        CUSTOM = CUSTOM.filter(function(x){ return x.c !== s.c; });
        delete BYCODE[s.c]; state.active.delete(s.c);
      });
      afterCustomChange();
      log('自定义清零', CUSTOM.length);
      log('只剩内置分组', tags().length);
      log('页头点数', txt(id('ptn')));
      out.push('');
      out.push('JS错误: ' + (errs.length ? errs.join(' | ') : '无'));
    }catch(e){
      out.push('✗ 探针中断：' + (e && e.message || e));
      out.push('JS错误: ' + (errs.length ? errs.join(' | ') : '无'));
    }
    fetch('/__tag__', { method: 'POST', body: out.join('\n') + '\n' });
  })();
})();
</script>
'''


class Quiet(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def do_POST(self):
        n = int(self.headers.get('Content-Length') or 0)
        RESULT.append(self.rfile.read(n).decode('utf-8', 'ignore'))
        self.send_response(200)
        self.send_header('Content-Length', '2')
        self.end_headers()
        self.wfile.write(b'ok')


def serve():
    h = functools.partial(Quiet, directory=SITE)
    socketserver.TCPServer.allow_reuse_address = True
    httpd = socketserver.TCPServer(('127.0.0.1', PORT), h)
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    return httpd


def run_probe(url, timeout=180):
    del RESULT[:]
    cmd = [EDGE, '--headless=new', '--no-sandbox', '--disable-gpu', '--disable-dev-shm-usage',
           '--window-size=1512,1600', url]
    p = subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    t0 = time.time()
    while time.time() - t0 < timeout:
        if RESULT:
            break
        time.sleep(0.25)
    time.sleep(0.4)
    try:
        p.terminate(); p.wait(timeout=8)
    except Exception:
        try: p.kill()
        except Exception: pass
    return RESULT[0] if RESULT else '（%.0f 秒内没收到探针回传）' % timeout


def main():
    shutil.rmtree(SITE, ignore_errors=True)
    os.makedirs(SITE, exist_ok=True)
    size, rows = make_xlsx()
    print('=' * 76)
    print('「商品标签分组 + 管理面板批量删除」E2E')
    print('=' * 76)
    print('测试用 xlsx：%s（%d 字节 · %d 行 · 3 个数值列 + 1 备注列）' % (XLSX_NAME, size, rows))

    src = io.open(SRC_PAGE, encoding='utf-8').read()
    # ★ 本机模式：把注入的云端配置掏空 —— 保证探针一条记录都不写进生产表
    stripped, cnt = re.subn(r'window\.WB_CFG\s*=\s*\{[^}]*\};', 'window.WB_CFG = null;', src)
    assert cnt == 1, '未能掏空 WB_CFG（匹配 %d 次）' % cnt
    # 进页面先把本机缓存清掉，保证每次跑都是干净起点（重复跑不会互相污染）
    i = src.find('<body')
    j = src.find('>', i)
    pre = ('<script>try{localStorage.removeItem("wb_c100_custom_v1");'
           'localStorage.removeItem("wb_c100_bover_v1");}catch(e){}</script>')
    page = stripped[:j + 1] + pre + stripped[j + 1:]
    io.open(os.path.join(SITE, PAGE), 'w', encoding='utf-8').write(page)
    io.open(os.path.join(SITE, PROBE_PAGE), 'w', encoding='utf-8').write(
        page.replace('</body>', DBG.replace('@XLSX@', XLSX_NAME) + '</body>', 1))

    srv = serve()
    res = run_probe('http://localhost:%d/%s' % (PORT, PROBE_PAGE), timeout=260)
    print(res)
    print('=' * 76)
    srv.shutdown()
    shutil.rmtree(SITE, ignore_errors=True)
    print('（已清理探针临时目录 %s）' % os.path.relpath(SITE, HERE))
    bad = ('✗' in res) or ('JS错误: 无' not in res)
    print('\n结论：' + ('✗ 有用例失败' if bad else '✓ 全部通过，零 JS 错误'))
    return 1 if bad else 0


if __name__ == '__main__':
    sys.exit(main())
