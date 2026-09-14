# -*- coding: utf-8 -*-
"""把 PPTX 里的文字 / 表格 / 原生图表数据完整抽出来，落 JSON。

原生图表的关键：shape.chart 里有 categories 与每条 series 的 name+values，
不用靠肉眼读图，可以直接还原成 ECharts 的 option。
"""
import io
import json
import os
import re

from pptx import Presentation
from pptx.util import Emu

OUT_DIR = r'C:\Users\50522\WorkBuddy\2026-09-10-14-50-38\build\src'


def walk(shapes, out, depth=0):
    for sh in shapes:
        try:
            if sh.shape_type is not None and 'GROUP' in str(sh.shape_type):
                walk(sh.shapes, out, depth + 1)
                continue
        except Exception:
            pass
        rec = {'name': getattr(sh, 'name', ''), 'depth': depth}
        try:
            rec['type'] = str(sh.shape_type)
        except Exception:
            rec['type'] = ''
        try:
            rec['pos'] = [int(sh.left or 0), int(sh.top or 0), int(sh.width or 0), int(sh.height or 0)]
        except Exception:
            rec['pos'] = None

        if sh.has_text_frame:
            t = sh.text_frame.text
            if t and t.strip():
                rec['text'] = t
        if getattr(sh, 'has_table', False) and sh.has_table:
            tb = sh.table
            rows = []
            for r in tb.rows:
                rows.append([c.text.replace('\n', ' ').strip() for c in r.cells])
            rec['table'] = rows
        if getattr(sh, 'has_chart', False) and sh.has_chart:
            ch = sh.chart
            info = {'chart_type': str(ch.chart_type)}
            try:
                info['has_title'] = ch.has_title
                if ch.has_title:
                    info['title'] = ch.chart_title.text_frame.text
            except Exception:
                pass
            cats = None
            try:
                pl = ch.plots[0]
                cats = [str(c) for c in (pl.categories or [])]
                info['plot_type'] = type(pl).__name__
            except Exception as e:
                info['cats_err'] = str(e)[:60]
            info['categories'] = cats
            sers = []
            for s in ch.series:
                try:
                    vals = list(s.values)
                except Exception:
                    vals = []
                vals = [None if v is None else (float(v) if isinstance(v, (int, float)) else str(v))
                        for v in vals]
                sers.append({'name': s.name, 'values': vals})
            info['series'] = sers
            rec['chart'] = info
        if 'picture' in rec['type'].lower():
            pass
        if len(rec) > 2:
            out.append(rec)


SRC = [
    (r'C:\Users\50522\OneDrive\工作研报周报\2026.07.27研究报告（鲁西化工更新&）\鲁西化工研究报告20260727.pptx', '000830', '鲁西化工'),
    (r'C:\Users\50522\OneDrive\工作研报周报\2026.07.06研究报告(四川美丰更新）\四川美丰研究报告（一）20250317.pptx', '000731', '四川美丰'),
    (r'C:\Users\50522\OneDrive\工作研报周报\2026.3.30研究报告（恒邦股份更新）\恒邦股份研究报告（更新）2026.3.30.pptx', '002237', '恒邦股份'),
]

os.makedirs(OUT_DIR, exist_ok=True)

for path, code, tag in SRC:
    prs = Presentation(path)
    slides = []
    for i, sl in enumerate(prs.slides, 1):
        items = []
        walk(sl.shapes, items)
        items.sort(key=lambda x: (x['pos'][1] if x['pos'] else 0, x['pos'][0] if x['pos'] else 0))
        slides.append({'page': i, 'items': items})
    doc = {'code': code, 'name': tag, 'src': os.path.basename(path), 'slides': slides}
    p = os.path.join(OUT_DIR, 'pptx_%s.json' % code)
    io.open(p, 'w', encoding='utf-8').write(json.dumps(doc, ensure_ascii=False, indent=1))

    n_txt = sum(1 for s in slides for it in s['items'] if 'text' in it)
    n_tbl = sum(1 for s in slides for it in s['items'] if 'table' in it)
    n_cht = sum(1 for s in slides for it in s['items'] if 'chart' in it)
    print('%-8s %-6s  页=%-3d 文本框=%-3d 表格=%-2d 原生图表=%-2d  → %s (%.0f KB)' % (
        code, tag, len(slides), n_txt, n_tbl, n_cht, os.path.basename(p), os.path.getsize(p) / 1024))
